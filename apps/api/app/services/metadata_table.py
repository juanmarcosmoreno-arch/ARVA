"""Sample metadata tables (CSV/TSV/TXT or .xlsx): reading, column mapping, validation and safe export.

Pure functions (no database, no network). Every cell is read as text (str) or None; nothing is coerced
to numbers, and dates are only accepted in explicit, unambiguous forms.
"""
from __future__ import annotations

import csv
import datetime as dt
import io
import json
import re
import threading
import unicodedata
import warnings as pywarnings
import zipfile
import zlib
from collections.abc import Iterable, Iterator
from contextlib import contextmanager
from dataclasses import dataclass, field
from pathlib import PurePath
from typing import Any
from xml.etree.ElementTree import ParseError, iterparse

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.utils.exceptions import CellCoordinatesException
from openpyxl.xml.constants import SHARED_STRINGS

from app.core.messages import render

XLSX_MAX_UNCOMPRESSED = 100 * 1024 * 1024
XLSX_MAX_MEMBERS = 2_000
# Only for parts above 1 MiB: small XML parts legitimately compress very well.
XLSX_MAX_RATIO = 250
# XML parts that openpyxl loads whole (shared strings, styles, workbook…) before any row is read.
XLSX_MAX_PART_BYTES = 16 * 1024 * 1024
XLSX_MAX_SHARED_STRINGS = 250_000
# Structural bounds checked in one streaming pass before openpyxl reads the sheet: empty or style-only
# <row>/<c> elements cost work too, and openpyxl pads every row to the widest cell it is asked for.
XLSX_MAX_ROW_ELEMENTS = 200_000
XLSX_MAX_CELL_ELEMENTS = 3_000_000
XLSX_MAX_GRID_CELLS = 4_000_000
EXCEL_MAX_ROW = 1_048_576
# Shared strings let a small workbook repeat one long text in every cell.
MAX_TOTAL_CHARS = 20_000_000
HEADER_SCAN_ROWS = 20
CSV_SNIFF_CHARS = 64 * 1024
DELIMITERS = ",;\t|"
CSV_EXTENSIONS = {".csv", ".tsv", ".txt"}
SAVE_AS = "save the sheet as .xlsx or CSV UTF-8"
UNSUPPORTED_EXTENSIONS = {
    ".xls": f"Legacy Excel 97-2003 .xls files are not supported; {SAVE_AS}.",
    ".xlsm": f"Macro-enabled workbooks (.xlsm) are not accepted; {SAVE_AS}.",
    ".xlsb": f"Binary workbooks (.xlsb) are not supported; {SAVE_AS}.",
    ".ods": f"OpenDocument spreadsheets (.ods) are not supported; {SAVE_AS}.",
    ".numbers": f"Apple Numbers files are not supported; {SAVE_AS}.",
}
OLE2_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"

ARVA_FIELDS = ("sample_code", "organism", "expected_species", "sample_type", "host", "geographic_origin",
               "collection_date", "sequencing_technology", "sequencing_run", "researcher", "institution",
               "description", "notes")
TARGETS = (*ARVA_FIELDS, "custom", "ignore")
FIELD_MAX_LEN = {"sample_code": 255, "organism": 255, "expected_species": 255, "sample_type": 100, "host": 255,
                 "geographic_origin": 255, "sequencing_run": 255, "researcher": 255, "institution": 255}

_CONTROL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
_ZERO_PAD_FMT = re.compile(r"^0+$")
_SCI_NOTATION = re.compile(r"[+-]?\d+(?:[.,]\d+)?[eE][+-]?\d+", re.ASCII)
_PLAIN_NUMBER = re.compile(r"[+-]?(?:\d+(?:[.,]\d+)?|[.,]\d+)(?:[eE][+-]?\d+)?", re.ASCII)
FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r", "\n", "＝", "＋", "－", "＠")
# warnings.catch_warnings() is process-global, so workbooks are read one at a time.
_XLSX_LOCK = threading.Lock()


# English text of every message in an import preview. The interface translates them by code (i18n keys
# metadata.msg.<code>, with the same {placeholders}) and shows this text for codes it does not know.
MESSAGES = {
    "header_detected": "The header was detected on row {row}; row(s) {rows} above it were ignored. Choose another "
                       "header row if this is wrong.",
    "column_unnamed": "Column {column} has data but no header; named '{name}'.",
    "header_duplicate": "Duplicate header '{header}' in column {column}; renamed '{name}'.",
    "row_extra_fields": "the row has {fields} fields but the header has {width}; extra values {values} (an unquoted "
                        "delimiter inside a field?)",
    "control_characters": "column '{column}' contains control characters; remove them",
    "csv_cp1252": "The file is not valid UTF-8, so it was read as Windows-1252 (typical of 'CSV' exports from Excel on "
                  "Windows). Check accented characters, or re-save it as 'CSV UTF-8'.",
    "csv_single_column": "No delimiter was found in the header line, so the file was read as a single column.",
    "csv_delimiter_guessed": "The delimiter could not be detected reliably (rows have different numbers of fields); "
                             "using {delimiter}, the most frequent one in the header line.",
    "csv_header_unprotected": "Header: the apostrophe that spreadsheet exports put before = + - @ to block formulas "
                              "was removed ({columns}).",
    "csv_cell_unprotected": "the apostrophe that spreadsheet exports put before = + - @ to block formulas was "
                            "removed in {columns}",
    "xlsx_error_value": "cell {cell}: Excel error value {value} imported as empty",
    "xlsx_number_padded": "cell {cell}: the number {value} has the display format {format}; imported as '{padded}' "
                          "as Excel shows it (format codes as Text in Excel to avoid this)",
    "xlsx_large_number": "cell {cell}: large number {value}; Excel keeps only 15 significant digits, so digits may "
                         "already be lost",
    "xlsx_duration": "cell {cell}: the duration {value} was imported as text",
    "xlsx_sheets": "The workbook has {count} sheets ({sheets}); the first one, '{sheet}', was read. Choose another "
                   "sheet if needed.",
    "xlsx_1904": "The workbook uses the 1904 date system; dates were converted accordingly.",
    "xlsx_formulas": "{count, plural, one {# formula cell} other {# formula cells}} ({cells}): the values stored by "
                     "the program that last saved the file were imported; formulas are never calculated.",
    "xlsx_formulas_uncached": "{count, plural, one {# formula cell ({cells}) has no stored value} other "
                              "{# formula cells ({cells}) have no stored value}} (the file was saved by a program "
                              "that does not calculate formulas) and {count, plural, one {was imported as empty} "
                              "other {were imported as empty}}. Open and re-save the file in Excel or LibreOffice, or "
                              "replace the formulas with values.",
    "xlsx_merged": "Merged cells ({ranges}): only the top-left cell of each range has a value; the other cells are "
                   "imported as empty, not copied.",
    "openpyxl": "openpyxl: {message}",
    "mapping_same_field": "Columns '{first}' and '{second}' both look like {field}; '{first}' is used and '{second}' "
                          "is kept as a custom column.",
    "technology_several": "technology {value} names several platforms ({platforms}); set to 'hybrid'",
    "technology_unsupported": "technology {value} is a platform ARVA has no value for (illumina, nanopore, pacbio, "
                              "hybrid); it is kept as text in the custom metadata",
    "technology_unrecognised": "unrecognised technology {value}; it is kept as text in the custom metadata",
    "date_invalid": "{value} is not a valid date (no such day, month or year)",
    "date_time_dropped": "the time of day in {value} was dropped (collection_date is a date)",
    "date_partial": "{value} is a partial date (year or year-month); collection_date needs a full date, so the text "
                    "is kept in the custom metadata",
    "date_no_month": "{value}: neither {first} nor {second} can be a month",
    "date_not_mdy": "{value} cannot be month/day/year (month {month})",
    "date_not_dmy": "{value} cannot be day/month/year (month {month})",
    "date_ambiguous": "{value} is ambiguous (day/month or month/day); use YYYY-MM-DD or choose the date order",
    "date_two_digit_year": "{value} has a two-digit year; use a four-digit year",
    "date_serial": "{value} looks like an Excel date serial number stored as a number; format the column as a date "
                   "or write YYYY-MM-DD",
    "date_unrecognised": "unrecognised date {value}; use YYYY-MM-DD",
    "date_future": "collection_date {date} is in the future",
    "date_before_1900": "collection_date {date} is before 1900",
    "no_sample_code_column": "No column is mapped to sample_code (for example 'sample_code', 'Sample ID' or 'ID "
                             "muestra'); map one column to it.",
    "date_order_from_header": "Date order {order} was taken from the header '{column}'.",
    "dates_mixed": "Column '{column}' mixes day-first (rows {dmy}) and month-first (rows {mdy}) dates; check the "
                   "file.",
    "dates_one_order": "Column '{column}': rows {rows} can only be {order}. Ambiguous dates are not assumed to follow "
                       "them; choose the date order {order} to confirm.",
    "sample_code_numbers": "The sample_code column '{column}' holds numbers in Excel (rows {rows}); leading zeros or "
                           "digits may have been lost. Format the column as Text in Excel.",
    "field_too_long": "{field} is {length} characters; the limit is {limit}",
    "row_no_sample_code_column": "no column is mapped to sample_code",
    "sample_code_empty": "sample_code is empty",
    "sample_code_whitespace": "sample_code contains a tab or line break",
    "sample_code_float": "sample_code {value} was stored in Excel as a non-integer number, so the original code "
                         "cannot be recovered; store sample codes as Text",
    "sample_code_scientific": "sample_code {value} looks like a number in scientific notation (possible spreadsheet "
                              "damage); check the original code",
    "sample_code_formula": "sample_code {value} starts with a spreadsheet formula character; it is stored as text and "
                           "neutralised on export",
    "sample_code_repeated": "sample_code {value} also appears on row {row}",
    "sample_code_repeated_case": "sample_code {value} also appears on row {row} as '{first}' (IDs that differ only in "
                                 "case are one sample)",
    "sample_code_ambiguous": "sample_code {value} matches several existing samples when case is ignored ({names}); "
                             "use the exact sample ID",
    "exists_create_only": "sample {sample} already exists and the mode is 'create only'",
}


class Msg(str):
    """A message's English text that also carries its code and parameters, so the interface can translate it."""

    code: str
    params: dict[str, str | int]

    def __new__(cls, code: str, **params: str | int) -> Msg:
        text = super().__new__(cls, render(MESSAGES[code], params))
        text.code, text.params = code, params
        return text


def message_json(message: str) -> dict[str, Any]:
    """{code, params, text}; code is null for free text without a translation."""
    return {"code": getattr(message, "code", None), "params": getattr(message, "params", {}), "text": str(message)}


def _rows(numbers: Iterable[int]) -> str:
    return ", ".join(str(n) for n in numbers)


class MetadataImportError(ValueError):
    """The file cannot be read at all; nothing is imported."""


class MetadataTooLargeError(MetadataImportError):
    pass


@dataclass(frozen=True)
class Limits:
    max_rows: int = 10_000
    max_columns: int = 256
    max_cell_chars: int = 10_000


def normalize_label(text: str) -> str:
    """Case- and accent-insensitive key: NFKD, combining marks dropped, casefolded, non-alphanumerics → space."""
    text = unicodedata.normalize("NFKD", str(text))
    text = "".join(ch for ch in text if not unicodedata.combining(ch)).casefold()
    return re.sub(r"[^0-9a-z]+", " ", text).strip()


def _clean_text(value: str) -> str | None:
    value = unicodedata.normalize("NFC", value).strip()
    return value or None


def _too_long(where: str, length: int, limits: Limits) -> MetadataTooLargeError:
    return MetadataTooLargeError(f"{where}: {length:,} characters; the limit is {limits.max_cell_chars:,}")


def _count_chars(cells: list[str | None], total: int) -> int:
    total += sum(len(c) for c in cells if c is not None)
    if total > MAX_TOTAL_CHARS:
        raise MetadataTooLargeError(f"the table holds more than {MAX_TOTAL_CHARS:,} characters of text")
    return total


# --- Header detection and table shaping (shared by CSV and XLSX) ----------------------------------

def _finalize_table(rows_in: list[tuple[int, list[str | None]]], header_row: int | None, limits: Limits,
                    warnings: list[str], label: str, strict_width: bool) -> tuple[list[str], list[dict], int, list[int]]:
    """rows_in: (row number, cells) of the non-empty rows only."""
    if not rows_in:
        raise MetadataImportError(f"{label}: no data found")
    if header_row is not None:
        found = [cells for n, cells in rows_in if n == header_row]
        if not found:
            raise MetadataImportError(f"{label}: row {header_row} is empty or does not exist, so it cannot be the header")
        hdr_n, hdr_cells = header_row, found[0]
    else:
        hdr_n, hdr_cells = _detect_header(rows_in[:HEADER_SCAN_ROWS])
        skipped = [n for n, _ in rows_in if n < hdr_n]
        if skipped:
            warnings.append(Msg("header_detected", row=hdr_n, rows=_rows(skipped)))
    for cell in hdr_cells:
        if cell is not None and _CONTROL.search(cell):
            raise MetadataImportError(f"{label}: header {cell!r} contains control characters; remove them")

    data_rows = [(n, cells) for n, cells in rows_in if n > hdr_n]
    width = max(len(c) for _, c in rows_in)
    hdr_cells = list(hdr_cells) + [None] * (width - len(hdr_cells))
    filled = {i for i, v in enumerate(hdr_cells) if v is not None}
    for _, cells in data_rows:
        filled.update(i for i, v in enumerate(cells) if v is not None)
    used = sorted(filled)

    columns: list[str] = []
    seen_keys: set[str] = set()
    for i in used:
        name = hdr_cells[i]
        if name is None:
            name = f"column_{get_column_letter(i + 1)}"
            warnings.append(Msg("column_unnamed", column=get_column_letter(i + 1), name=name))
        key = normalize_label(name) or name
        if key in seen_keys or name in columns:
            k = 2
            while f"{name} ({k})" in columns:
                k += 1
            renamed = f"{name} ({k})"
            warnings.append(Msg("header_duplicate", header=name, column=get_column_letter(i + 1), name=renamed))
            name = renamed
        seen_keys.add(key)
        columns.append(name)

    hdr_width = len(hdr_cells)
    while hdr_width and hdr_cells[hdr_width - 1] is None:
        hdr_width -= 1
    rows: list[dict] = []
    for n, cells in data_rows:
        errors: list[str] = []
        if strict_width and len(cells) > hdr_width:
            extra = [c for c in cells[hdr_width:] if c is not None]
            if extra:
                errors.append(Msg("row_extra_fields", fields=len(cells), width=hdr_width, values=repr(extra)))
        values = {col: (cells[i] if i < len(cells) else None) for col, i in zip(columns, used, strict=True)}
        for col, val in values.items():
            if val is not None and _CONTROL.search(val):
                errors.append(Msg("control_characters", column=col))
        rows.append({"row": n, "values": values, "errors": errors, "warnings": []})
        if len(rows) > limits.max_rows:
            raise MetadataTooLargeError(f"{label}: more than {limits.max_rows:,} data rows")
    return columns, rows, hdr_n, used


def _detect_header(rows: list[tuple[int, list[str | None]]]) -> tuple[int, list[str | None]]:
    """First row with a recognised sample_code header, else the first with two or more cells, else the first."""
    for n, cells in rows:
        if any(c is not None and FIELD_BY_SYNONYM.get(_header_key(c)) == "sample_code" for c in cells):
            return n, cells
    for n, cells in rows:
        if sum(c is not None for c in cells) >= 2:
            return n, cells
    return rows[0]


def _trim(cells: list[str | None]) -> list[str | None]:
    while cells and cells[-1] is None:
        cells.pop()
    return cells


# --- CSV ----------------------------------------------------------------------------------------

def _decode_csv(data: bytes, warnings: list[str]) -> tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "utf-8-sig"
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16"), "utf-16"
        except UnicodeDecodeError as exc:
            raise MetadataImportError(f"the file has a UTF-16 byte-order mark but is not valid UTF-16: {exc}") from exc
    try:
        return data.decode("utf-8"), "utf-8"
    except UnicodeDecodeError as exc:
        first_error = exc
    try:
        text = data.decode("cp1252")
    except UnicodeDecodeError as exc:
        raise MetadataImportError(f"the file is neither valid UTF-8 ({first_error}) nor Windows-1252 ({exc}); "
                                  "re-save it as 'CSV UTF-8'") from exc
    warnings.append(Msg("csv_cp1252"))
    return text, "cp1252"


def _choose_delimiter(text: str, filename: str, warnings: list[str]) -> str:
    if PurePath(filename).suffix.lower() == ".tsv":
        return "\t"
    sample = text[:CSV_SNIFF_CHARS]
    if len(text) > CSV_SNIFF_CHARS and "\n" in sample:
        sample = sample[: sample.rfind("\n") + 1]
    try:
        # Restricting the candidates matters: unrestricted, the sniffer can return a letter.
        return csv.Sniffer().sniff(sample, delimiters=DELIMITERS).delimiter
    except csv.Error:
        pass
    first_line = next((ln for ln in sample.splitlines() if ln.strip()), "")
    counts = {d: first_line.count(d) for d in DELIMITERS}
    best = max(counts.values())
    if best == 0:
        warnings.append(Msg("csv_single_column"))
        return ","
    winners = [d for d, c in counts.items() if c == best]
    if len(winners) > 1:
        raise MetadataImportError("cannot determine the delimiter: the header line contains "
                                  f"{' and '.join(repr(w) for w in winners)} equally often; choose the delimiter")
    warnings.append(Msg("csv_delimiter_guessed", delimiter=repr(winners[0])))
    return winners[0]


def _parse_csv(data: bytes, filename: str, delimiter: str | None, header_row: int | None, limits: Limits) -> dict:
    warnings: list[str] = []
    if b"\x00" in data and not data.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise MetadataImportError("the file contains NUL bytes, so it is not a text CSV file")
    text, encoding = _decode_csv(data, warnings)
    if "\x00" in text:
        raise MetadataImportError("the file contains NUL characters after decoding")
    # Excel's "sep=;" first line; Excel hides it, so row numbers start after it.
    sep = re.match(r"sep=([,;\t|])\r?\n", text)
    if sep:
        delimiter = delimiter or sep.group(1)
        text = text[sep.end():]
    used_delimiter = delimiter or _choose_delimiter(text, filename, warnings)

    reader = csv.reader(io.StringIO(text, newline=""), delimiter=used_delimiter, quotechar='"', doublequote=True,
                        strict=False)
    rows_in: list[tuple[int, list[str | None]]] = []
    unprotected: dict[int, list[int]] = {}
    record = total = 0
    try:
        for fields in reader:
            record += 1
            for j, f in enumerate(fields):
                if len(f) > limits.max_cell_chars:
                    raise _too_long(f"row {record}, field {j + 1}", len(f), limits)
            cells = _trim([_clean_text(f) for f in fields])
            for j, cell in enumerate(cells):
                if cell is not None and is_neutralised_cell(cell):
                    cells[j] = _clean_text(cell[1:])
                    unprotected.setdefault(record, []).append(j)
            if not cells:
                continue
            total = _count_chars(cells, total)
            if len(cells) > limits.max_columns:
                raise MetadataTooLargeError(f"CSV row {record} has {len(cells)} columns; the limit is {limits.max_columns}")
            rows_in.append((record, cells))
            if len(rows_in) > limits.max_rows + HEADER_SCAN_ROWS:
                raise MetadataTooLargeError(f"CSV: more than {limits.max_rows:,} data rows")
    except csv.Error as exc:
        if "field larger than field limit" in str(exc):
            raise MetadataTooLargeError(f"CSV row {record + 1}: a field is longer than {limits.max_cell_chars:,} "
                                        "characters") from exc
        raise MetadataImportError(f"CSV parse error near row {record + 1}: {exc}") from exc

    columns, rows, hdr, used = _finalize_table(rows_in, header_row, limits, warnings, "CSV", strict_width=True)
    if unprotected:
        by_row = {r["row"]: r for r in rows}
        names = dict(zip(used, columns, strict=True))
        for n, idx in unprotected.items():
            where = ", ".join(f"'{names[j]}'" for j in idx if j in names)
            if n == hdr:
                warnings.append(Msg("csv_header_unprotected", columns=where))
            elif n in by_row and where:
                by_row[n]["warnings"].append(Msg("csv_cell_unprotected", columns=where))
    return {"kind": "csv", "encoding": encoding, "delimiter": used_delimiter, "sheet": None, "sheets": [],
            "header_row": hdr, "columns": columns, "rows": rows, "warnings": warnings, "numeric_cells": {}}


# --- XLSX ---------------------------------------------------------------------------------------

# What zipfile, zlib, expat and openpyxl's coordinate parser raise on damaged files (RuntimeError: encrypted part).
_DAMAGED_FILE_ERRORS = (zipfile.BadZipFile, zlib.error, ParseError, UnicodeDecodeError, ValueError, KeyError, OSError,
                        EOFError, RuntimeError, NotImplementedError, CellCoordinatesException)


@contextmanager
def _damaged(what: str) -> Iterator[None]:
    try:
        yield
    except MetadataImportError:
        raise
    except _DAMAGED_FILE_ERRORS as exc:
        raise MetadataImportError(f"{what}: {type(exc).__name__}: {exc}") from exc


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _check_zip(data: bytes) -> None:
    with _damaged("not a valid .xlsx (zip) file"), zipfile.ZipFile(io.BytesIO(data)) as zf:
        infos = zf.infolist()
        if len(infos) > XLSX_MAX_MEMBERS:
            raise MetadataTooLargeError(f"the .xlsx has {len(infos)} parts; the limit is {XLSX_MAX_MEMBERS}")
        total = 0
        for info in infos:
            total += info.file_size
            if info.file_size > 1024 * 1024 and info.compress_size and info.file_size / info.compress_size > XLSX_MAX_RATIO:
                raise MetadataTooLargeError(
                    f"the .xlsx part {info.filename} expands {info.file_size / info.compress_size:.0f}x (limit "
                    f"{XLSX_MAX_RATIO}x); refusing to open it (possible zip bomb)")
        if total > XLSX_MAX_UNCOMPRESSED:
            raise MetadataTooLargeError(f"the .xlsx expands to {total:,} bytes; the limit is {XLSX_MAX_UNCOMPRESSED:,}")
        for info in infos:
            if (info.filename.endswith((".xml", ".rels")) and not info.filename.startswith("xl/worksheets/")
                    and info.file_size > XLSX_MAX_PART_BYTES):
                raise MetadataTooLargeError(f"the .xlsx part {info.filename} is {info.file_size:,} bytes; the limit is "
                                            f"{XLSX_MAX_PART_BYTES:,}")
        names = {i.filename for i in infos}
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise MetadataImportError("the zip file is not an Excel workbook ([Content_Types].xml or xl/workbook.xml "
                                      "is missing)")
    with _damaged("cannot open the workbook"), zipfile.ZipFile(io.BytesIO(data)) as zf:
        _check_shared_strings(zf)


def _check_shared_strings(zf: zipfile.ZipFile) -> None:
    """openpyxl loads every shared string into memory when the workbook is opened, so count them first."""
    part = None
    with zf.open("[Content_Types].xml") as fh:
        for _event, elem in iterparse(fh, events=("end",)):
            if _local(elem.tag) == "Override" and elem.get("ContentType") == SHARED_STRINGS:
                part = (elem.get("PartName") or "").lstrip("/")
                break
    if not part or part not in zf.namelist():
        return
    if zf.getinfo(part).file_size > XLSX_MAX_PART_BYTES:
        raise MetadataTooLargeError(f"the shared-strings part {part} is {zf.getinfo(part).file_size:,} bytes; the "
                                    f"limit is {XLSX_MAX_PART_BYTES:,}")
    count = 0
    root = None
    with zf.open(part) as fh:
        for event, elem in iterparse(fh, events=("start", "end")):
            if root is None:
                root = elem
            elif event == "end" and _local(elem.tag) == "si":
                count += 1
                if count > XLSX_MAX_SHARED_STRINGS:
                    raise MetadataTooLargeError(f"the workbook holds more than {XLSX_MAX_SHARED_STRINGS:,} distinct "
                                                "texts; save the metadata sheet on its own as .xlsx or CSV UTF-8")
                root.clear()


@dataclass
class _SheetScan:
    last_row: int = 0  # last row and column holding a value, an inline string or a formula
    last_col: int = 0
    merged: list[str] = field(default_factory=list)
    merged_count: int = 0
    formulas: list[str] = field(default_factory=list)
    formula_count: int = 0
    uncached: list[str] = field(default_factory=list)
    uncached_count: int = 0


def _row_number(text: str) -> int:
    value = float(text)  # openpyxl accepts '7' and '7.0'
    if not value.is_integer():
        raise ValueError(f"{text!r} is not a row number")
    return int(value)


def _scan_sheet_xml(data: bytes, sheet_path: str, label: str, limits: Limits) -> _SheetScan:
    """One streaming pass before openpyxl reads the sheet: bounds its size and finds merged ranges and formula
    cells, which read-only worksheets do not expose."""
    scan = _SheetScan()
    rows = cells = data_rows = row_n = col_n = 0
    in_row = row_has_data = False
    sheet_data = None
    with _damaged(f"cannot read {label}"), zipfile.ZipFile(io.BytesIO(data)) as zf, zf.open(sheet_path) as fh:
        for event, elem in iterparse(fh, events=("start", "end")):
            tag = _local(elem.tag)
            if event == "start":
                if tag == "sheetData":
                    sheet_data = elem
                elif tag == "row":
                    rows += 1
                    if rows > XLSX_MAX_ROW_ELEMENTS:
                        raise MetadataTooLargeError(f"{label} has more than {XLSX_MAX_ROW_ELEMENTS:,} rows, counting "
                                                    "empty formatted ones")
                    r = elem.get("r")
                    number = _row_number(r) if r is not None else row_n + 1
                    if number <= row_n or number > EXCEL_MAX_ROW:
                        raise MetadataImportError(f"{label}: row number {r} is out of order or out of range; open "
                                                  "and re-save the file in Excel or LibreOffice")
                    row_n, col_n, in_row, row_has_data = number, 0, True, False
                continue
            if tag == "c" and in_row:
                cells += 1
                if cells > XLSX_MAX_CELL_ELEMENTS:
                    raise MetadataTooLargeError(f"{label} has more than {XLSX_MAX_CELL_ELEMENTS:,} cells, counting "
                                                "empty formatted ones")
                ref = elem.get("r")
                col_n = coordinate_to_tuple(ref)[1] if ref else col_n + 1
                kids = {_local(child.tag): child for child in elem}
                coord = _coord(row_n, col_n - 1)
                if "f" in kids:
                    scan.formula_count += 1
                    if len(scan.formulas) < 10:
                        scan.formulas.append(coord)
                    if not ("v" in kids and kids["v"].text):
                        scan.uncached_count += 1
                        if len(scan.uncached) < 10:
                            scan.uncached.append(coord)
                if kids.keys() & {"v", "is", "f"}:
                    if col_n > limits.max_columns:
                        raise MetadataTooLargeError(f"{label}: cell {coord} holds a value in column {col_n}; the limit "
                                                    f"is {limits.max_columns} columns")
                    row_has_data = True
                    scan.last_col = max(scan.last_col, col_n)
                elem.clear()
            elif tag == "row":
                in_row = False
                if row_has_data:
                    data_rows += 1
                    scan.last_row = row_n
                    if data_rows > limits.max_rows + HEADER_SCAN_ROWS:
                        raise MetadataTooLargeError(f"{label}: more than {limits.max_rows:,} data rows")
                elem.clear()
                if sheet_data is not None:
                    sheet_data.clear()
            elif tag == "mergeCell":
                scan.merged_count += 1
                if len(scan.merged) < 10:
                    scan.merged.append(elem.get("ref") or "?")
                elem.clear()
    if scan.last_row * scan.last_col > XLSX_MAX_GRID_CELLS:
        raise MetadataTooLargeError(
            f"{label}: values reach row {scan.last_row:,} and column {scan.last_col} "
            f"({_coord(scan.last_row, scan.last_col - 1)}); delete stray cells far below or to the right of the table")
    return scan


def _coord(row: int, col: int) -> str:
    return f"{get_column_letter(col + 1)}{row}"


def _xlsx_cell_to_text(cell: Any, r: int, c: int, warnings: list[tuple[int, str]]) -> tuple[str | None, bool]:
    """(text, came from a number); warnings collects (row, message)."""
    value = cell.value
    if value is None:
        return None, False

    def warn(code: str, **params: str | int) -> None:
        warnings.append((r, Msg(code, cell=_coord(r, c), **params)))

    if getattr(cell, "data_type", None) == "e":
        warn("xlsx_error_value", value=repr(value))
        return None, False
    if isinstance(value, bool):
        return ("TRUE" if value else "FALSE"), False
    if isinstance(value, int):
        text = str(value)
        fmt = getattr(cell, "number_format", "General") or "General"
        if _ZERO_PAD_FMT.match(fmt) and len(fmt) > len(text) and value >= 0:
            padded = text.zfill(len(fmt))
            warn("xlsx_number_padded", value=value, format=repr(fmt), padded=padded)
            text = padded
        return text, True
    if isinstance(value, float):
        if value.is_integer() and abs(value) < 2**53:
            return str(int(value)), True
        if abs(value) >= 1e15:
            warn("xlsx_large_number", value=repr(value))
        return repr(value), True
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat(), False
        return value.isoformat(timespec="seconds"), False
    if isinstance(value, dt.date | dt.time):
        return value.isoformat(), False
    if isinstance(value, dt.timedelta):
        warn("xlsx_duration", value=str(value))
        return str(value), False
    return _clean_text(str(value)), False


def _parse_xlsx(data: bytes, sheet: str | None, header_row: int | None, limits: Limits) -> dict:
    _check_zip(data)
    with _XLSX_LOCK, pywarnings.catch_warnings(record=True) as caught:
        pywarnings.simplefilter("always")
        result = _read_xlsx(data, sheet, header_row, limits)
    forwarded: list[str] = []
    for w in caught:
        message = Msg("openpyxl", message=str(w.message))
        if "openpyxl" in w.filename and message not in forwarded:
            forwarded.append(message)
    result["warnings"].extend(forwarded)
    return result


def _read_xlsx(data: bytes, sheet: str | None, header_row: int | None, limits: Limits) -> dict:
    warnings: list[str] = []
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # BadZipFile, KeyError, XML ParseError (entity amplification) and others
        raise MetadataImportError(f"cannot open the workbook: {type(exc).__name__}: {exc}") from exc
    cell_warnings: list[tuple[int, str]] = []
    numeric: dict[int, list[int]] = {}
    rows_in: list[tuple[int, list[str | None]]] = []
    try:
        sheets = [ws.title for ws in wb.worksheets]
        if not sheets:
            raise MetadataImportError("the workbook has no worksheets")
        if sheet is None:
            ws = wb.worksheets[0]
            if len(sheets) > 1:
                warnings.append(Msg("xlsx_sheets", count=len(sheets), sheets=", ".join(sheets), sheet=ws.title))
        elif sheet in sheets:
            ws = wb[sheet]
        else:
            raise MetadataImportError(f"sheet {sheet!r} not found; the workbook has: {', '.join(sheets)}")
        if wb.epoch.year == 1904:
            warnings.append(Msg("xlsx_1904"))
        title = ws.title
        label = f"sheet '{title}'"
        sheet_path = getattr(ws, "_worksheet_path", None)  # private attribute; openpyxl is pinned
        if not sheet_path:
            raise MetadataImportError(f"cannot locate the XML of {label}")
        scan = _scan_sheet_xml(data, sheet_path, label, limits)
        # The stored <dimension> can be wrong; the bounds come from the scan instead.
        ws.reset_dimensions()
        total = 0
        try:
            rows_iter = ws.iter_rows(max_row=scan.last_row, max_col=scan.last_col) if scan.last_row else ()
            for r_idx, row in enumerate(rows_iter, start=1):
                cells: list[str | None] = []
                for c_idx, cell in enumerate(row):
                    text, is_number = _xlsx_cell_to_text(cell, r_idx, c_idx, cell_warnings)
                    if text is not None and len(text) > limits.max_cell_chars:
                        raise _too_long(f"cell {_coord(r_idx, c_idx)}", len(text), limits)
                    cells.append(text)
                    if is_number:
                        numeric.setdefault(c_idx, []).append(r_idx)
                if not _trim(cells):
                    continue
                total = _count_chars(cells, total)
                if len(cells) > limits.max_columns:
                    raise MetadataTooLargeError(f"{label} row {r_idx} has {len(cells)} columns; the limit "
                                                f"is {limits.max_columns}")
                rows_in.append((r_idx, cells))
                if len(rows_in) > limits.max_rows + HEADER_SCAN_ROWS:
                    raise MetadataTooLargeError(f"{label}: more than {limits.max_rows:,} data rows")
        except MetadataImportError:
            raise
        except Exception as exc:
            raise MetadataImportError(f"cannot read {label}: {type(exc).__name__}: {exc}") from exc
    finally:
        wb.close()

    if scan.formula_count:
        warnings.append(Msg("xlsx_formulas", count=scan.formula_count, cells=", ".join(scan.formulas)))
    if scan.uncached_count:
        warnings.append(Msg("xlsx_formulas_uncached", count=scan.uncached_count, cells=", ".join(scan.uncached)))
    if scan.merged_count:
        warnings.append(Msg("xlsx_merged", ranges=", ".join(scan.merged)))

    columns, rows, hdr, used = _finalize_table(rows_in, header_row, limits, warnings, label, strict_width=False)
    by_row = {r["row"]: r for r in rows}
    for r, message in cell_warnings:
        if r in by_row:
            by_row[r]["warnings"].append(message)
    numeric_cells = {}
    for col, idx in zip(columns, used, strict=True):
        found = {r for r in numeric.get(idx, []) if r > hdr}
        if found:
            numeric_cells[col] = found
    return {"kind": "xlsx", "encoding": None, "delimiter": None, "sheet": title, "sheets": sheets, "header_row": hdr,
            "columns": columns, "rows": rows, "warnings": warnings, "numeric_cells": numeric_cells}


def parse_metadata_table(data: bytes, filename: str, *, sheet: str | None = None, delimiter: str | None = None,
                         header_row: int | None = None, limits: Limits = Limits()) -> dict:
    """Read a CSV/TSV/TXT or .xlsx table into columns and rows of str|None values; raises MetadataImportError."""
    ext = PurePath(filename).suffix.lower()
    if ext in UNSUPPORTED_EXTENSIONS:
        raise MetadataImportError(UNSUPPORTED_EXTENSIONS[ext])
    if delimiter is not None and delimiter not in tuple(DELIMITERS):
        raise MetadataImportError("the delimiter must be one of: comma, semicolon, tab or |")
    if not data:
        raise MetadataImportError("the file is empty")
    if data.startswith(OLE2_MAGIC):
        raise MetadataImportError(UNSUPPORTED_EXTENSIONS[".xls"] + " (the content is a legacy OLE2 workbook)")
    if ext == ".xlsx":
        if not data.startswith(ZIP_MAGIC):
            raise MetadataImportError("the file has the .xlsx extension but is not an .xlsx (zip) file")
        return _parse_xlsx(data, sheet, header_row, limits)
    if ext in CSV_EXTENSIONS:
        if data.startswith(ZIP_MAGIC):
            raise MetadataImportError(f"the file has the {ext} extension but is a zip archive (a renamed .xlsx?)")
        return _parse_csv(data, filename, delimiter, header_row, limits)
    raise MetadataImportError(f"unsupported file type {ext or '(no extension)'}; use .csv, .tsv, .txt or .xlsx")


# --- Column mapping -----------------------------------------------------------------------------

SYNONYMS: dict[str, list[str]] = {
    "sample_code": ["sample_code", "sample code", "sample", "sample id", "sampleid", "sample name", "sample number",
                    "isolate id", "codigo", "codigo de muestra", "codigo muestra", "cod muestra", "muestra",
                    "id muestra", "id de muestra", "id_muestra", "identificador", "identificador de muestra",
                    "nombre de muestra", "numero de muestra", "id aislado"],
    "organism": ["organism", "organismo", "microorganism", "microorganismo"],
    "expected_species": ["expected species", "expected_species", "species", "especie", "especie esperada",
                         "suspected species", "especie sospechada", "presumptive species", "especie presuntiva"],
    "sample_type": ["sample type", "sample_type", "tipo de muestra", "tipo muestra", "specimen", "specimen type",
                    "tipo de especimen", "isolation source", "fuente de aislamiento"],
    "host": ["host", "hospedador", "hospedero", "huesped", "host species", "especie hospedadora"],
    "geographic_origin": ["geographic origin", "geographic_origin", "geographical origin", "origin", "origen",
                          "origen geografico", "country", "pais", "location", "ubicacion", "localizacion",
                          "procedencia", "geo loc name", "geo_loc_name"],
    "collection_date": ["collection date", "collection_date", "date", "fecha", "fecha de muestreo", "fecha muestreo",
                        "fecha de recoleccion", "fecha de colecta", "fecha de toma", "fecha de toma de muestra",
                        "sampling date", "date of collection"],
    "sequencing_technology": ["sequencing technology", "sequencing_technology", "technology", "tecnologia",
                              "tecnologia de secuenciacion", "platform", "sequencing platform", "plataforma",
                              "plataforma de secuenciacion", "sequencer", "secuenciador", "instrument",
                              "instrumento"],
    "sequencing_run": ["sequencing run", "sequencing_run", "run", "run id", "corrida", "corrida de secuenciacion",
                       "run de secuenciacion"],
    "researcher": ["researcher", "investigator", "investigador", "investigadora", "investigador responsable",
                   "responsable", "principal investigator"],
    "institution": ["institution", "institucion", "organization", "organisation", "organizacion", "centro",
                    "center", "centre"],
    "description": ["description", "descripcion"],
    "notes": ["notes", "note", "notas", "nota", "observaciones", "comentarios", "comments", "remarks"],
}
_PAREN = re.compile(r"\([^)]*\)|\[[^\]]*\]")


def _header_key(header: str) -> str:
    return normalize_label(_PAREN.sub(" ", header))


FIELD_BY_SYNONYM: dict[str, str] = {normalize_label(s): field for field, syns in SYNONYMS.items() for s in syns}


def suggest_field(header: str) -> str:
    """Exact synonym match after normalisation (no fuzzy matching); anything else is a custom column."""
    return FIELD_BY_SYNONYM.get(_header_key(header), "custom")


_DATE_HINTS = {
    "DMY": re.compile(r"\b(dd|d)\W?(mm|m)\W?(aaaa|yyyy|aa|yy)\b", re.I),
    "MDY": re.compile(r"\b(mm|m)\W?(dd|d)\W?(aaaa|yyyy|aa|yy)\b", re.I),
    "YMD": re.compile(r"\b(aaaa|yyyy)\W?(mm|m)\W?(dd|d)\b", re.I),
}


def date_order_hint(header: str) -> str | None:
    """'Fecha (dd/mm/aaaa)' → 'DMY'. Only an explicit format written in the header counts."""
    for order, rx in _DATE_HINTS.items():
        if rx.search(header):
            return order
    return None


def map_columns(columns: Iterable[str], overrides: dict[str, str] | None = None) -> dict:
    """{header: ARVA field | 'custom' | 'ignore'}; user choices win over suggestions."""
    columns = list(columns)
    overrides = overrides or {}
    unknown = [k for k in overrides if k not in columns]
    if unknown:
        raise MetadataImportError(f"the mapping names column(s) that are not in the file: {unknown[:5]}")
    invalid = sorted({v for v in overrides.values() if v not in TARGETS})
    if invalid:
        raise MetadataImportError(f"unknown target(s) in the mapping: {invalid[:5]}")
    taken: dict[str, str] = {}
    for col in columns:
        field = overrides.get(col)
        if field in ARVA_FIELDS:
            if field in taken:
                raise MetadataImportError(f"columns '{taken[field]}' and '{col}' are both mapped to {field}; map only one")
            taken[field] = col
    mapping: dict[str, str] = {}
    warnings: list[str] = []
    for col in columns:
        if col in overrides:
            mapping[col] = overrides[col]
            continue
        field = suggest_field(col)
        if field in taken:
            warnings.append(Msg("mapping_same_field", first=taken[field], second=col, field=field))
            field = "custom"
        elif field != "custom":
            taken[field] = col
        mapping[col] = field
    return {"mapping": mapping, "warnings": warnings}


# --- Values -------------------------------------------------------------------------------------

TECH_PATTERNS = {
    "illumina": re.compile(r"\b(illumina|miseq|hiseq|nextseq|novaseq|iseq|miniseq|solexa|genome analyzer)"),
    "nanopore": re.compile(r"\b(nanopore|ont|minion|gridion|promethion|flongle|oxford nanopore)\b"),
    "pacbio": re.compile(r"\b(pacbio|pacific biosciences|hifi|sequel|revio|smrt|rs ii|rsii|onso)\b"),
}
TECH_HYBRID = {"hybrid", "hibrido", "hibrida", "hybrid assembly", "short long", "long short"}
MISSING_TOKENS = {"unknown", "desconocido", "desconocida", "na", "n a", "nd", "n d", "none", "missing",
                  "not collected", "not provided", "not applicable", "sin dato", "sin datos", "no disponible"}
TECH_UNSUPPORTED = re.compile(r"\b(ion torrent|iontorrent|ion s5|ion pgm|ion proton|mgi|dnbseq|bgiseq|bgi|"
                              r"aviti|element biosciences|sanger|capillary|454|solid|ultima)\b")


def normalize_technology(value: str | None) -> tuple[str | None, str | None]:
    """(ARVA technology or None when the cell carries no information, warning)."""
    if value is None:
        return None, None
    key = normalize_label(value)
    if not key or key in MISSING_TOKENS:
        return None, None
    if key in {"illumina", "nanopore", "pacbio", "hybrid"}:
        return key, None
    if key in TECH_HYBRID:
        return "hybrid", None
    families = [fam for fam, rx in TECH_PATTERNS.items() if rx.search(key)]
    if len(families) == 1:
        return families[0], None
    if len(families) > 1:
        return "hybrid", Msg("technology_several", value=repr(value), platforms=", ".join(families))
    if TECH_UNSUPPORTED.search(key):
        return "unknown", Msg("technology_unsupported", value=repr(value))
    return "unknown", Msg("technology_unrecognised", value=repr(value))


MONTHS = {
    "january": 1, "jan": 1, "february": 2, "feb": 2, "march": 3, "mar": 3, "april": 4, "apr": 4, "may": 5,
    "june": 6, "jun": 6, "july": 7, "jul": 7, "august": 8, "aug": 8, "september": 9, "sep": 9, "sept": 9,
    "october": 10, "oct": 10, "november": 11, "nov": 11, "december": 12, "dec": 12,
    "enero": 1, "ene": 1, "febrero": 2, "marzo": 3, "abril": 4, "abr": 4, "mayo": 5, "junio": 6, "julio": 7,
    "agosto": 8, "ago": 8, "septiembre": 9, "setiembre": 9, "set": 9, "octubre": 10, "noviembre": 11,
    "diciembre": 12, "dic": 12,
}
DATE_MISSING = MISSING_TOKENS | {"sin fecha", "no date"}
# Statuses that block the row; 'partial' only warns (the text is kept, collection_date stays unset).
BLOCKING_DATE_STATUSES = ("ambiguous", "invalid", "two_digit_year", "unrecognised", "serial")


@dataclass
class DateResult:
    value: dt.date | None
    status: str  # ok | empty | ambiguous | invalid | partial | two_digit_year | unrecognised | serial
    message: str | None = None
    evidence: str | None = None  # 'DMY' or 'MDY' when the value itself proves the order


def parse_date(text: str | None, date_order: str | None = None) -> DateResult:
    """Explicit dates only: ISO, YYYY/MM/DD, D/M/YYYY or M/D/YYYY when unambiguous or date_order is given,
    and English or Spanish month names. Day/month order is never guessed."""
    if date_order not in (None, "DMY", "MDY"):
        raise ValueError("date_order must be None, 'DMY' or 'MDY'")
    if text is None:
        return DateResult(None, "empty")
    s = text.strip()
    if not s or normalize_label(s) in DATE_MISSING:
        return DateResult(None, "empty")

    def make(y: Any, m: Any, d: Any, evidence: str | None = None) -> DateResult:
        try:
            return DateResult(dt.date(int(y), int(m), int(d)), "ok", None, evidence)
        except ValueError:
            return DateResult(None, "invalid", Msg("date_invalid", value=repr(text)))

    m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[T ](\d{2}):(\d{2})(?::(\d{2})(?:\.\d+)?)?)?", s, re.ASCII)
    if m:
        res = make(*m.group(1, 2, 3))
        if res.value and m.group(4) and (m.group(4), m.group(5), m.group(6) or "00") != ("00", "00", "00"):
            res.message = Msg("date_time_dropped", value=repr(text))
        return res
    m = re.fullmatch(r"(\d{4})[/.](\d{1,2})[/.](\d{1,2})", s, re.ASCII)
    if m:
        return make(*m.group(1, 2, 3))
    if re.fullmatch(r"\d{4}(-\d{2})?", s, re.ASCII):
        return DateResult(None, "partial", Msg("date_partial", value=repr(text)))
    m = re.fullmatch(r"(\d{1,2})([/.\-])(\d{1,2})\2(\d{4})", s, re.ASCII)
    if m:
        a, b, y = int(m.group(1)), int(m.group(3)), m.group(4)
        if a > 12 and b > 12:
            return DateResult(None, "invalid", Msg("date_no_month", value=repr(text), first=a, second=b))
        if a == b:
            return make(y, a, b)
        if a > 12:
            if date_order == "MDY":
                return DateResult(None, "invalid", Msg("date_not_mdy", value=repr(text), month=a))
            return make(y, b, a, evidence="DMY")
        if b > 12:
            if date_order == "DMY":
                return DateResult(None, "invalid", Msg("date_not_dmy", value=repr(text), month=b))
            return make(y, a, b, evidence="MDY")
        if date_order == "DMY":
            return make(y, b, a)
        if date_order == "MDY":
            return make(y, a, b)
        return DateResult(None, "ambiguous", Msg("date_ambiguous", value=repr(text)))
    if re.fullmatch(r"\d{1,2}[/.\-]\d{1,2}[/.\-]\d{2}", s, re.ASCII):
        return DateResult(None, "two_digit_year", Msg("date_two_digit_year", value=repr(text)))
    tokens = [t for t in normalize_label(s).split() if t not in {"de", "del", "of", "the"}]
    if len(tokens) == 3:
        month = [t for t in tokens if t in MONTHS]
        years = [t for t in tokens if t.isascii() and t.isdigit() and len(t) == 4]
        days = [t for t in tokens if t.isascii() and t.isdigit() and len(t) <= 2]
        if len(month) == 1 and len(years) == 1 and len(days) == 1:
            return make(years[0], MONTHS[month[0]], days[0])
    if re.fullmatch(r"\d{5}(\.\d+)?", s, re.ASCII) and 20000 <= float(s) <= 80000:
        return DateResult(None, "serial", Msg("date_serial", value=repr(text)))
    return DateResult(None, "unrecognised", Msg("date_unrecognised", value=repr(text)))


# --- Rows → ARVA sample fields ------------------------------------------------------------------

def _unique_key(custom: dict, values: dict, key: str) -> str:
    base, i = key, 2
    while key in custom or key in values:
        key = f"{base} {i}"
        i += 1
    return key


def build_samples(parsed: dict, *, overrides: dict[str, str] | None = None, date_order: str | None = None,
                  today: dt.date | None = None) -> dict:
    """Map, validate and normalise parse_metadata_table() rows. Only non-empty cells appear in a row's data."""
    today = today or dt.date.today()
    mapped = map_columns(parsed["columns"], overrides)
    mapping = mapped["mapping"]
    warnings = list(parsed["warnings"]) + mapped["warnings"]
    errors: list[str] = []
    by_field = {f: c for c, f in mapping.items() if f in ARVA_FIELDS}
    code_col = by_field.get("sample_code")
    if code_col is None:
        errors.append(Msg("no_sample_code_column"))

    date_col = by_field.get("collection_date")
    date_order_source = "user" if date_order else None
    if date_order is None and date_col and date_order_hint(date_col) in ("DMY", "MDY"):
        date_order, date_order_source = date_order_hint(date_col), "header"
        warnings.append(Msg("date_order_from_header", order=date_order, column=date_col))
    ambiguous_dates = 0
    if date_col:
        evidence: dict[str, list[int]] = {"DMY": [], "MDY": []}
        for r in parsed["rows"]:
            res = parse_date(r["values"].get(date_col), date_order)
            ambiguous_dates += res.status == "ambiguous"
            if res.evidence and date_order is None:
                evidence[res.evidence].append(r["row"])
        if evidence["DMY"] and evidence["MDY"]:
            warnings.append(Msg("dates_mixed", column=date_col, dmy=_rows(evidence["DMY"][:5]),
                                mdy=_rows(evidence["MDY"][:5])))
        elif ambiguous_dates and (evidence["DMY"] or evidence["MDY"]):
            order = "DMY" if evidence["DMY"] else "MDY"
            warnings.append(Msg("dates_one_order", column=date_col, rows=_rows((evidence["DMY"] or evidence["MDY"])[:5]),
                                order=order))

    numeric_code_rows = parsed["numeric_cells"].get(code_col, set()) if code_col else set()
    if numeric_code_rows:
        warnings.append(Msg("sample_code_numbers", column=code_col, rows=_rows(sorted(numeric_code_rows)[:10])))

    out_rows = []
    for r in parsed["rows"]:
        n, values = r["row"], r["values"]
        row_errors, row_warnings = list(r["errors"]), list(r["warnings"])
        data: dict[str, str] = {}
        custom: dict[str, str] = {}
        for col, val in values.items():
            target = mapping[col]
            if target == "ignore" or val is None:
                continue
            if target == "custom":
                custom[col] = val
            elif target == "collection_date":
                res = parse_date(val, date_order)
                if res.value:
                    data[target] = res.value.isoformat()
                    if res.message:
                        row_warnings.append(res.message)
                    if res.value > today:
                        row_warnings.append(Msg("date_future", date=res.value.isoformat()))
                    if res.value.year < 1900:
                        row_warnings.append(Msg("date_before_1900", date=res.value.isoformat()))
                elif res.status != "empty":
                    custom[_unique_key(custom, values, f"{col} (original)")] = val
                    (row_errors if res.status in BLOCKING_DATE_STATUSES else row_warnings).append(res.message)
            elif target == "sequencing_technology":
                tech, message = normalize_technology(val)
                if message:
                    row_warnings.append(message)
                    custom[_unique_key(custom, values, f"{col} (original)")] = val
                if tech is not None and tech != "unknown":
                    data[target] = tech
            elif target in FIELD_MAX_LEN and len(val) > FIELD_MAX_LEN[target]:
                row_errors.append(Msg("field_too_long", field=target, length=len(val), limit=FIELD_MAX_LEN[target]))
            else:
                data[target] = val
        code = data.get("sample_code")
        if code_col is None:
            row_errors.append(Msg("row_no_sample_code_column"))
        elif values.get(code_col) is None:
            row_errors.append(Msg("sample_code_empty"))
        if code:
            if re.search(r"[\t\r\n]", code):
                row_errors.append(Msg("sample_code_whitespace"))
            if n in numeric_code_rows and re.search(r"[.eE]", code):
                row_errors.append(Msg("sample_code_float", value=repr(code)))
            elif _SCI_NOTATION.fullmatch(code):
                row_warnings.append(Msg("sample_code_scientific", value=repr(code)))
            if code.startswith(FORMULA_PREFIXES):
                row_warnings.append(Msg("sample_code_formula", value=repr(code)))
        out_rows.append({"row": n, "sample_code": code, "data": data, "custom": custom,
                         "errors": row_errors, "warnings": row_warnings})

    columns = [{"header": c, "field": mapping[c], "suggested_field": suggest_field(c)} for c in parsed["columns"]]
    return {"columns": columns, "mapping": mapping, "rows": out_rows, "warnings": warnings, "errors": errors,
            "date_order": date_order, "date_order_source": date_order_source, "ambiguous_dates": ambiguous_dates}


def _as_text(value: Any) -> str | None:
    if value is None or isinstance(value, str):
        return value
    return json.dumps(value, ensure_ascii=False)


def plan_rows(built: dict, existing: list[dict]) -> list[dict]:
    """Decide create / update / unchanged / error per row against the project's samples.

    existing: [{"id", "sample_code", <ARVA fields as str|None>, "custom_metadata": dict|None}]. A row matches a
    sample by exact sample_code, else by a unique case-insensitive match; several case-insensitive matches are an error.
    """
    exact = {s["sample_code"]: s for s in existing}
    folded: dict[str, list[dict]] = {}
    for s in existing:
        folded.setdefault(s["sample_code"].strip().casefold(), []).append(s)
    seen: dict[str, tuple[int, str]] = {}
    planned = []
    for r in built["rows"]:
        errors = list(r["errors"])
        code = r["sample_code"]
        target, matched_by = None, None
        if code:
            key = code.casefold()
            if key in seen:
                first_row, first_code = seen[key]
                if first_code == code:
                    errors.append(Msg("sample_code_repeated", value=repr(code), row=first_row))
                else:
                    errors.append(Msg("sample_code_repeated_case", value=repr(code), row=first_row, first=first_code))
            else:
                seen[key] = (r["row"], code)
            if code in exact:
                target, matched_by = exact[code], "exact"
            elif len(folded.get(key, [])) == 1:
                target, matched_by = folded[key][0], "case_insensitive"
            elif folded.get(key):
                names = ", ".join(sorted(s["sample_code"] for s in folded[key]))
                errors.append(Msg("sample_code_ambiguous", value=repr(code), names=names))
        if not code and not errors:
            errors.append(Msg("sample_code_empty"))
        changes: list[dict] = []
        if not errors:
            old_custom = (target or {}).get("custom_metadata") or {}
            for field, new in r["data"].items():
                old = target.get(field) if target else None
                if field != "sample_code" and old != new:
                    changes.append({"field": field, "key": None, "old": old, "new": new})
            for key, new in r["custom"].items():
                old = _as_text(old_custom.get(key))
                if old != new:
                    changes.append({"field": "custom_metadata", "key": key, "old": old, "new": new})
        if errors:
            action = "error"
        elif target is None:
            action = "create"
            changes.insert(0, {"field": "sample_code", "key": None, "old": None, "new": code})
        else:
            action = "update" if changes else "unchanged"
        planned.append({**r, "errors": errors, "action": action, "changes": changes,
                        "sample_id": target["id"] if target else None,
                        "matched_sample_code": target["sample_code"] if target else None, "matched_by": matched_by})
    return planned


# --- Export (formula-injection safe) ------------------------------------------------------------

def _formula_like(s: str) -> bool:
    return s.startswith(FORMULA_PREFIXES) and not (s[0] in "+-" and bool(_PLAIN_NUMBER.fullmatch(s)))


def sanitize_spreadsheet_cell(value: Any) -> str:
    """OWASP CSV-injection mitigation: prefix ' to text starting with = + - @ TAB CR LF (or full-width = + - @).
    Plain numbers such as -34.6 are left alone, since a sign is not a formula there. Apostrophes already in
    front of such text get one more, so that import can always remove exactly one (is_neutralised_cell)."""
    if value is None:
        return ""
    s = str(value)
    return "'" + s if _formula_like(s.lstrip("'")) else s


def is_neutralised_cell(text: str) -> bool:
    """True exactly for the cells sanitize_spreadsheet_cell() changed, so CSV export → import is lossless."""
    return text.startswith("'") and _formula_like(text.lstrip("'"))


def export_table(rows: list[dict]) -> tuple[list[str], list[list[str | None]]]:
    """Header (ARVA fields, then custom_metadata keys in first-seen order) and text rows."""
    custom_keys: list[str] = []
    for row in rows:
        for key in row.get("custom_metadata") or {}:
            if key not in custom_keys:
                custom_keys.append(key)
    table = []
    for row in rows:
        custom = row.get("custom_metadata") or {}
        table.append([_as_text(row.get(f)) for f in ARVA_FIELDS] + [_as_text(custom.get(k)) for k in custom_keys])
    return [*ARVA_FIELDS, *custom_keys], table


def spreadsheet_csv(header: list[str], table: list[list[Any]]) -> bytes:
    """UTF-8 with BOM (so Excel detects the encoding), CRLF, every field quoted and neutralised."""
    buf = io.StringIO(newline="")
    writer = csv.writer(buf, quoting=csv.QUOTE_ALL, lineterminator="\r\n")
    writer.writerow([sanitize_spreadsheet_cell(h) for h in header])
    for values in table:
        writer.writerow([sanitize_spreadsheet_cell(v) for v in values])
    return ("﻿" + buf.getvalue()).encode("utf-8")


def export_csv(rows: list[dict]) -> bytes:
    return spreadsheet_csv(*export_table(rows))


def write_xlsx_text_cell(ws: Any, row: int, column: int, value: str | None) -> None:
    """A text cell: openpyxl would otherwise turn '=…' into a formula and '#N/A' into an error value."""
    if value is None:
        return
    cell = ws.cell(row=row, column=column, value=ILLEGAL_CHARACTERS_RE.sub("�", value))
    cell.data_type = "s"


def export_xlsx(rows: list[dict], sheet_title: str = "Samples") -> bytes:
    header, table = export_table(rows)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title
    # Text format keeps codes such as 001 intact when they are edited in Excel.
    ws.column_dimensions["A"].number_format = "@"
    ws.freeze_panes = "A2"
    for c, name in enumerate(header, start=1):
        write_xlsx_text_cell(ws, 1, c, name)
    for r, values in enumerate(table, start=2):
        for c, value in enumerate(values, start=1):
            write_xlsx_text_cell(ws, r, c, value)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
