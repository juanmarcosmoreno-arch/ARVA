"""Sample metadata import (preview / apply), history, export and template endpoints."""
import csv
import datetime as dt
import hashlib
import io
import json
import re
from pathlib import Path

import openpyxl
import pytest
from sqlalchemy import func, select, text

from app.api.metadata_limit import FORM_OVERHEAD_BYTES, MetadataUploadLimit
from app.core.config import get_settings
from app.db.session import SessionLocal, engine
from app.models import MetadataImport, Sample
from app.services import metadata_import

API = "/api/v1"
FIX = Path(__file__).parent / "fixtures" / "metadata"
FASTQ = b"@r1\nACGT\n+\nIIII\n@r2\nGGGCCC\n+\n555555\n"


def make_project(client, name="Metadata survey"):
    r = client.post(f"{API}/projects", json={"name": name})
    assert r.status_code == 201, r.text
    return r.json()


def make_sample(client, project_id, code, **extra):
    r = client.post(f"{API}/projects/{project_id}/samples", json={"sample_code": code, **extra})
    assert r.status_code == 201, r.text
    return r.json()


def samples_by_code(client, project_id):
    return {s["sample_code"]: s for s in client.get(f"{API}/projects/{project_id}/samples").json()}


def post_table(client, project_id, content: bytes, name="meta.csv", action="preview", **form):
    data = {k: (json.dumps(v) if k == "mapping" else str(v)) for k, v in form.items() if v is not None}
    return client.post(f"{API}/projects/{project_id}/metadata/{action}",
                       files={"file": (name, io.BytesIO(content), "application/octet-stream")}, data=data)


def apply_table(client, project_id, content: bytes, name="meta.csv", **form):
    form.setdefault("expected_sha256", hashlib.sha256(content).hexdigest())
    return post_table(client, project_id, content, name, action="apply", **form)


def count_imports() -> int:
    with SessionLocal() as db:
        return db.scalar(select(func.count(MetadataImport.id)))


TABLE = (
    "Sample ID,Hospedador,País,Fecha de muestreo,Tecnología,MLST,Notas\n"
    "NEW-1,Bos taurus,Peru,2024-03-15,MiSeq,ST131,first\n"
    "s-old,,Chile,,,ST22,\n"
    "SAME,Homo sapiens,,,,,\n"
    ",orphan,,,,,\n"
)


def test_preview_shows_actions_and_writes_nothing(client):
    p = make_project(client)
    make_sample(client, p["id"], "S-OLD", host="Gallus gallus", sequencing_technology="illumina",
                custom_metadata={"MLST": "ST1", "lab": "A"})
    make_sample(client, p["id"], "SAME", host="Homo sapiens")
    r = post_table(client, p["id"], TABLE.encode())
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["file"]["kind"] == "csv" and body["file"]["delimiter"] == "," and body["file"]["header_row"] == 1
    assert body["file"]["sha256"] == hashlib.sha256(TABLE.encode()).hexdigest()
    cols = {c["header"]: c for c in body["columns"]}
    assert cols["Sample ID"]["field"] == "sample_code" and cols["Hospedador"]["field"] == "host"
    assert cols["MLST"]["field"] == "custom" and cols["MLST"]["suggested_field"] == "custom"
    assert "ignore" in body["fields"] and "sample_code" in body["fields"]
    assert body["counts"] == {"rows": 4, "create": 1, "update": 1, "unchanged": 1, "error": 1}
    rows = {r["row"]: r for r in body["rows"]}
    assert rows[2]["action"] == "create"
    assert {(c["field"], c["new"]) for c in rows[2]["changes"]} >= {
        ("sample_code", "NEW-1"), ("host", "Bos taurus"), ("collection_date", "2024-03-15"),
        ("sequencing_technology", "illumina"), ("notes", "first")}
    assert rows[3]["action"] == "update" and rows[3]["matched_by"] == "case_insensitive"
    assert rows[3]["matched_sample_code"] == "S-OLD"
    assert rows[3]["changes"] == [
        {"field": "geographic_origin", "key": None, "old": None, "new": "Chile"},
        {"field": "custom_metadata", "key": "MLST", "old": "ST1", "new": "ST22"},
    ]
    assert rows[4]["action"] == "unchanged"
    assert rows[5]["action"] == "error"
    assert rows[5]["errors"] == [{"code": "sample_code_empty", "params": {}, "text": "sample_code is empty"}]
    assert set(samples_by_code(client, p["id"])) == {"S-OLD", "SAME"}
    assert count_imports() == 0


def test_apply_creates_updates_without_clearing_and_records_provenance(client):
    p = make_project(client)
    old = make_sample(client, p["id"], "S-OLD", host="Gallus gallus", sequencing_technology="illumina",
                      collection_date="2020-01-02", notes="keep me", custom_metadata={"MLST": "ST1", "lab": "A"})
    make_sample(client, p["id"], "SAME", host="Homo sapiens")
    r = apply_table(client, p["id"], TABLE.encode(), name="../../evil/meta.csv")
    assert r.status_code == 201, r.text
    record, skipped = r.json()["record"], r.json()["skipped"]
    assert (record["created_count"], record["updated_count"], record["unchanged_count"], record["skipped_count"]) == (1, 1, 1, 1)
    assert record["filename"] == "meta.csv" and record["kind"] == "csv" and record["mode"] == "create_and_update"
    assert record["mapping"]["Sample ID"] == "sample_code" and record["mapping"]["MLST"] == "custom"
    assert record["sha256"] == hashlib.sha256(TABLE.encode()).hexdigest() and record["header_row"] == 1
    assert skipped == [{"row": 5, "sample_code": None,
                        "reasons": [{"code": "sample_code_empty", "params": {}, "text": "sample_code is empty"}]}]

    samples = samples_by_code(client, p["id"])
    assert set(samples) == {"NEW-1", "S-OLD", "SAME"}
    new = samples["NEW-1"]
    assert (new["host"], new["geographic_origin"], new["collection_date"], new["sequencing_technology"]) == (
        "Bos taurus", "Peru", "2024-03-15", "illumina")
    assert new["notes"] == "first" and new["custom_metadata"] == {"MLST": "ST131"}
    updated = samples["S-OLD"]
    assert updated["id"] == old["id"]
    # empty cells never clear: host, date, technology and notes keep their stored values
    assert (updated["host"], updated["collection_date"], updated["sequencing_technology"], updated["notes"]) == (
        "Gallus gallus", "2020-01-02", "illumina", "keep me")
    assert updated["geographic_origin"] == "Chile"
    assert updated["custom_metadata"] == {"MLST": "ST22", "lab": "A"}

    history = client.get(f"{API}/projects/{p['id']}/metadata/imports").json()
    assert [h["id"] for h in history] == [record["id"]]


def test_create_only_mode_skips_existing_samples(client):
    p = make_project(client)
    make_sample(client, p["id"], "S-OLD")
    make_sample(client, p["id"], "SAME", host="Homo sapiens")
    r = apply_table(client, p["id"], TABLE.encode(), mode="create_only")
    assert r.status_code == 201, r.text
    record = r.json()["record"]
    assert (record["created_count"], record["updated_count"], record["skipped_count"]) == (1, 0, 2)
    reasons = [reason for s in r.json()["skipped"] for reason in s["reasons"]]
    assert {"code": "exists_create_only", "params": {"sample": "'S-OLD'"},
            "text": "sample 'S-OLD' already exists and the mode is 'create only'"} in reasons
    assert samples_by_code(client, p["id"])["S-OLD"]["geographic_origin"] is None


def test_ambiguous_case_insensitive_match_is_a_row_error_and_skipped(client):
    p = make_project(client)
    make_sample(client, p["id"], "Ab-1")
    make_sample(client, p["id"], "AB-1")
    content = b"sample,host\nab-1,cow\nNEW,pig\n"
    rows = post_table(client, p["id"], content).json()["rows"]
    assert rows[0]["action"] == "error" and "matches several existing samples" in rows[0]["errors"][0]["text"]
    assert rows[0]["errors"][0]["code"] == "sample_code_ambiguous" and rows[0]["errors"][0]["params"]["value"] == "'ab-1'"
    r = apply_table(client, p["id"], content)
    assert r.json()["record"]["created_count"] == 1 and r.json()["skipped"][0]["row"] == 2
    assert all(s["host"] is None for c, s in samples_by_code(client, p["id"]).items() if c != "NEW")


def test_sha256_mismatch_is_refused(client):
    p = make_project(client)
    r = apply_table(client, p["id"], b"sample\nS1\n", expected_sha256="0" * 64)
    assert r.status_code == 409 and "Preview it again" in r.json()["detail"]
    assert apply_table(client, p["id"], b"sample\nS1\n", expected_sha256="nothex").status_code == 422
    assert samples_by_code(client, p["id"]) == {}


def test_apply_is_one_transaction(client, monkeypatch):
    p = make_project(client)
    make_sample(client, p["id"], "S-OLD", host="Gallus gallus")

    def broken(**_kw):
        raise RuntimeError("disk on fire")

    monkeypatch.setattr(metadata_import, "MetadataImport", broken)
    with pytest.raises(RuntimeError):
        apply_table(client, p["id"], TABLE.encode())
    samples = samples_by_code(client, p["id"])
    assert set(samples) == {"S-OLD"} and samples["S-OLD"]["geographic_origin"] is None
    assert count_imports() == 0


def test_nothing_to_apply_is_refused(client):
    p = make_project(client)
    make_sample(client, p["id"], "S1", host="cow")
    r = apply_table(client, p["id"], b"sample,host\nS1,cow\n,x\n")
    assert r.status_code == 400 and "Nothing to apply" in r.json()["detail"]
    assert count_imports() == 0


def test_mapping_overrides_custom_and_ignore(client):
    p = make_project(client)
    content = b"Cepa;Lugar;Secreto;MLST\nC1;Lima;x;ST5\n"
    mapping = {"Cepa": "sample_code", "Lugar": "geographic_origin", "Secreto": "ignore", "MLST": "custom"}
    body = post_table(client, p["id"], content, mapping=mapping).json()
    assert {c["header"]: c["field"] for c in body["columns"]} == mapping
    assert body["counts"]["create"] == 1 and not body["errors"]
    r = apply_table(client, p["id"], content, mapping=mapping)
    assert r.status_code == 201, r.text
    s = samples_by_code(client, p["id"])["C1"]
    assert s["geographic_origin"] == "Lima" and s["custom_metadata"] == {"MLST": "ST5"}

    no_code = post_table(client, p["id"], content).json()
    assert no_code["errors"] and no_code["counts"]["error"] == 1

    for bad, needle in (("{", "not valid JSON"), ("[1]", "JSON object")):
        r = client.post(f"{API}/projects/{p['id']}/metadata/preview",
                        files={"file": ("m.csv", io.BytesIO(content))}, data={"mapping": bad})
        assert r.status_code == 400 and needle in r.json()["detail"]
    r = post_table(client, p["id"], content, mapping={"Nope": "host"})
    assert r.status_code == 400 and "not in the file" in r.json()["detail"]
    r = post_table(client, p["id"], content, mapping={"Cepa": "host", "Lugar": "host"})
    assert r.status_code == 400 and "both mapped to host" in r.json()["detail"]


def test_ambiguous_dates_block_rows_until_a_date_order_is_chosen(client):
    p = make_project(client)
    content = b"sample;fecha\nA;15/03/2024\nB;03/05/2024\n"
    body = post_table(client, p["id"], content).json()
    assert body["ambiguous_dates"] == 1 and body["date_order"] is None
    assert [r["action"] for r in body["rows"]] == ["create", "error"]
    body = post_table(client, p["id"], content, date_order="DMY").json()
    assert body["ambiguous_dates"] == 0 and body["date_order"] == "DMY" and body["date_order_source"] == "user"
    r = apply_table(client, p["id"], content, date_order="DMY")
    assert r.json()["record"]["date_order"] == "DMY"
    assert samples_by_code(client, p["id"])["B"]["collection_date"] == "2024-05-03"
    assert post_table(client, p["id"], content, date_order="YMD").status_code == 422


def test_xlsx_preview_with_sheet_and_header_row(client):
    p = make_project(client)
    wb = openpyxl.Workbook()
    wb.active.title = "Notes"
    wb.active.append(["nothing here"])
    ws = wb.create_sheet("Muestras")
    ws.append(["Proyecto 2024"])
    ws.append(["ID muestra", "Organismo", "Fecha"])
    ws.append(["001", "E. coli", dt.date(2024, 3, 15)])
    buf = io.BytesIO()
    wb.save(buf)
    body = post_table(client, p["id"], buf.getvalue(), name="m.xlsx", sheet="Muestras").json()
    assert body["file"]["kind"] == "xlsx" and body["file"]["sheets"] == ["Notes", "Muestras"]
    assert body["file"]["sheet"] == "Muestras" and body["file"]["header_row"] == 2
    row = body["rows"][0]
    assert row["sample_code"] == "001" and row["row"] == 3
    assert {"field": "collection_date", "key": None, "old": None, "new": "2024-03-15"} in row["changes"]
    r = post_table(client, p["id"], buf.getvalue(), name="m.xlsx", sheet="Nope")
    assert r.status_code == 400 and "not found" in r.json()["detail"]
    assert post_table(client, p["id"], buf.getvalue(), name="m.xlsx", sheet="Muestras", header_row=0).status_code == 422


def test_fixture_files_preview(client):
    p = make_project(client)
    body = post_table(client, p["id"], (FIX / "lo_resaved_typed.xlsx").read_bytes(), name="lo.xlsx").json()
    assert body["counts"]["rows"] == 6 and body["counts"]["error"] >= 1
    assert any(w["code"] == "xlsx_formulas" for w in body["warnings"])
    body = post_table(client, p["id"], (FIX / "lo_export_cp1252_semicolon.csv").read_bytes()).json()
    assert body["file"]["encoding"] == "cp1252" and body["file"]["delimiter"] == ";"


def test_unsupported_and_broken_files_are_rejected(client):
    p = make_project(client)
    r = post_table(client, p["id"], b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100, name="old.xls")
    assert r.status_code == 400 and "save the sheet as .xlsx or CSV UTF-8" in r.json()["detail"]
    assert post_table(client, p["id"], b"a,b\n", name="fake.xlsx").status_code == 400
    assert post_table(client, p["id"], b"", name="m.csv").status_code == 400
    assert post_table(client, p["id"], b"a,b\n1,\x002\n").status_code == 400
    assert post_table(client, p["id"], b"sample\nS1\n", delimiter=";;").status_code == 422
    assert post_table(client, p["id"], b"sample\nS1\n", delimiter="x").status_code == 400
    assert post_table(client, "nope", b"sample\nS1\n").status_code == 404



def test_damaged_xlsx_zip_is_a_400_not_a_500(client):
    p = make_project(client)
    wb = openpyxl.Workbook()
    wb.active.append(["sample_code"])
    buf = io.BytesIO()
    wb.save(buf)
    data = bytearray(buf.getvalue())
    entry = data.find(b"PK\x01\x02")
    data[entry + 9] |= 0x08  # the central directory says the name is UTF-8; its bytes are not
    data[entry + 46:entry + 48] = b"\xc3\x28"
    for action in ("preview", "apply"):
        r = post_table(client, p["id"], bytes(data), name="bad.xlsx", action=action,
                       expected_sha256=hashlib.sha256(bytes(data)).hexdigest() if action == "apply" else None)
        assert r.status_code == 400 and "not a valid .xlsx (zip) file" in r.json()["detail"]


def test_apply_parses_the_file_before_locking_the_project(client, monkeypatch):
    p = make_project(client)
    real = metadata_import.mt.parse_metadata_table
    lockable = []

    def parse_and_probe(*args, **kwargs):
        with engine.connect() as other:
            other.execute(text("SELECT id FROM projects WHERE id = :id FOR UPDATE NOWAIT"), {"id": p["id"]})
            lockable.append(True)
            other.rollback()
        return real(*args, **kwargs)

    monkeypatch.setattr(metadata_import.mt, "parse_metadata_table", parse_and_probe)
    assert apply_table(client, p["id"], b"sample\nS1\n").status_code == 201
    assert lockable == [True]

def test_size_and_row_limits(client, monkeypatch):
    p = make_project(client)
    settings = get_settings()
    monkeypatch.setattr(settings, "metadata_max_bytes", 50)
    r = post_table(client, p["id"], b"sample,notes\n" + b"S1,x\n" * 20)
    assert r.status_code == 413 and "ARVA_METADATA_MAX_BYTES" in r.json()["detail"]
    r = post_table(client, p["id"], b"x" * (50 + FORM_OVERHEAD_BYTES + 1))
    assert r.status_code == 413
    monkeypatch.setattr(settings, "metadata_max_bytes", 10 * 1024 * 1024)
    monkeypatch.setattr(settings, "metadata_max_rows", 3)
    r = post_table(client, p["id"], b"sample\n" + b"".join(b"S%d\n" % i for i in range(30)))
    assert r.status_code == 413 and "more than 3 data rows" in r.json()["detail"]
    monkeypatch.setattr(settings, "metadata_max_columns", 2)
    assert post_table(client, p["id"], b"a,b,c\n1,2,3\n").status_code == 413


async def test_upload_limit_answers_before_reading_the_body():
    reached = []

    async def inner(scope, receive, send):
        reached.append(scope["path"])

    async def receive():
        raise AssertionError("the request body must not be read")

    sent = []

    async def send(message):
        sent.append(message)

    guard = MetadataUploadLimit(inner, "/api/v1")
    too_big = str(get_settings().metadata_max_bytes + FORM_OVERHEAD_BYTES + 1).encode()
    scope = {"type": "http", "method": "POST", "path": "/api/v1/projects/abc/metadata/preview",
             "headers": [(b"content-length", too_big)]}
    await guard(scope, receive, send)
    await guard({**scope, "path": "/api/v1/projects/abc/metadata/apply"}, receive, send)
    assert [m["status"] for m in sent if m["type"] == "http.response.start"] == [413, 413] and reached == []
    await guard({**scope, "headers": []}, receive, send)
    assert sent[-2]["status"] == 411
    await guard({**scope, "headers": [(b"content-length", b"100")]}, receive, send)
    await guard({**scope, "path": "/api/v1/projects/abc/metadata/imports"}, receive, send)
    assert reached == ["/api/v1/projects/abc/metadata/preview", "/api/v1/projects/abc/metadata/imports"]


EVIL_CODE = "=HYPERLINK(\"http://x\",\"y\")"
INJECTION = {"notes": "@SUM(1+1)*cmd|' /C calc'!A0", "host": "+34 600",
             "custom_metadata": {"lat": "-34.6", "=evil": "#N/A", "n": 3}}


def test_export_csv_is_injection_safe(client):
    p = make_project(client, name="Survey: 2024/25 <b>")
    make_sample(client, p["id"], EVIL_CODE, **INJECTION)
    make_sample(client, p["id"], "001", collection_date="2024-03-15", sequencing_technology="pacbio")
    r = client.get(f"{API}/projects/{p['id']}/metadata/export", params={"format": "csv"})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    assert re.fullmatch(r'attachment; filename="ARVA_Survey_2024_25_b_metadata_\d{8}T\d{6}Z\.csv"',
                        r.headers["content-disposition"])
    assert r.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"), newline="")))
    by_code = {row["sample_code"]: row for row in rows}
    evil = by_code["'" + EVIL_CODE]
    assert evil["notes"] == "'@SUM(1+1)*cmd|' /C calc'!A0" and evil["host"] == "'+34 600"
    assert evil["lat"] == "-34.6" and evil["'=evil"] == "#N/A" and evil["n"] == "3"
    assert by_code["001"]["collection_date"] == "2024-03-15" and by_code["001"]["sequencing_technology"] == "pacbio"
    assert list(rows[0])[:3] == ["sample_code", "organism", "expected_species"]


def test_csv_export_reimports_unchanged(client):
    p = make_project(client)
    make_sample(client, p["id"], "-NEG-CTRL", notes="-80 °C freezer", host="=lab", custom_metadata={"Catalase": "+", "Gram": "-"})
    make_sample(client, p["id"], EVIL_CODE, **INJECTION)
    make_sample(client, p["id"], "S3", geographic_origin="-34.6, -58.4", custom_metadata={"lat": "-34.6", "'ref": "'+x", "q": "''=y"})
    exported = client.get(f"{API}/projects/{p['id']}/metadata/export", params={"format": "csv"}).content
    assert b'"\'-NEG-CTRL"' in exported and b'"\'+"' in exported
    preview = post_table(client, p["id"], exported, name="export.csv").json()
    assert preview["counts"]["unchanged"] == 3 and preview["counts"]["rows"] == 3, preview["rows"]
    neg = next(r for r in preview["rows"] if r["sample_code"] == "-NEG-CTRL")
    assert any(w["code"] == "csv_cell_unprotected" and "'Catalase'" in w["params"]["columns"] for w in neg["warnings"])
    assert any(w["code"] == "csv_header_unprotected" for w in preview["warnings"])  # the custom key '=evil'
    assert apply_table(client, p["id"], exported, name="export.csv").status_code == 400  # nothing to apply


def test_export_filename_keeps_accented_letters_as_ascii(client):
    p = make_project(client, name='Proyecto Perú ñ "quoted" / Señal')
    r = client.get(f"{API}/projects/{p['id']}/metadata/export", params={"format": "xlsx"})
    assert re.fullmatch(r'attachment; filename="ARVA_Proyecto_Peru_n_quoted_Senal_metadata_\d{8}T\d{6}Z\.xlsx"',
                        r.headers["content-disposition"])


def test_export_xlsx_is_injection_safe(client):
    p = make_project(client)
    make_sample(client, p["id"], EVIL_CODE, **INJECTION)
    r = client.get(f"{API}/projects/{p['id']}/metadata/export", params={"format": "xlsx"})
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    assert r.headers["content-disposition"].endswith('.xlsx"')
    wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=False)
    ws = wb.active
    cells = [c for row in ws.iter_rows() for c in row if c.value is not None]
    assert cells and all(c.data_type == "s" for c in cells)
    values = {c.value for c in cells}
    assert {EVIL_CODE, INJECTION["notes"], "#N/A", "=evil", "-34.6", "3"} <= values
    assert post_table(client, p["id"], r.content, name="export.xlsx").json()["counts"]["unchanged"] == 1


def test_template_is_header_only(client):
    r = client.get(f"{API}/metadata/template", params={"format": "csv"})
    assert r.status_code == 200 and 'filename="ARVA_sample_metadata_template.csv"' in r.headers["content-disposition"]
    lines = r.content.decode("utf-8-sig").strip().split("\r\n")
    assert len(lines) == 1 and lines[0].startswith('"sample_code","organism"')
    r = client.get(f"{API}/metadata/template", params={"format": "xlsx"})
    ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
    assert ws.max_row == 1 and ws["A1"].value == "sample_code" and ws.column_dimensions["A"].number_format == "@"
    assert client.get(f"{API}/metadata/template", params={"format": "xls"}).status_code == 422


def test_report_csv_neutralises_text_but_not_numbers(client):
    p = make_project(client)
    s = make_sample(client, p["id"], "=cmd|' /C calc'!A0")
    client.post(f"{API}/samples/{s['id']}/files",
                files=[("files", ("@evil.fastq", io.BytesIO(FASTQ), "application/octet-stream"))])
    client.post(f"{API}/samples/{s['id']}/analyses/sequence-stats", json={})
    r = client.get(f"{API}/samples/{s['id']}/report", params={"format": "csv"})
    row = next(csv.DictReader(io.StringIO(r.text)))
    assert row["sample_code"] == "'=cmd|' /C calc'!A0" and row["filename"] == "'@evil.fastq"
    assert row["reads"] == "2" and row["derived_from"] == ""


def test_project_delete_removes_import_history(client):
    p = make_project(client)
    apply_table(client, p["id"], b"sample\nS1\n")
    assert count_imports() == 1
    client.delete(f"{API}/projects/{p['id']}")
    assert count_imports() == 0
    with SessionLocal() as db:
        assert db.scalar(select(func.count(Sample.id))) == 0
