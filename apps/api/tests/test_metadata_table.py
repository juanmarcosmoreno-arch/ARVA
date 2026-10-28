"""Pure functions of app.services.metadata_table (ported from the research prototype's tests)."""
import csv
import datetime as dt
import io
import json
import re
import zipfile
from pathlib import Path

import openpyxl
import pytest
from openpyxl.utils import get_column_letter

from app.services import metadata_table as mt

FIX = Path(__file__).parent / "fixtures" / "metadata"
TODAY = dt.date(2026, 10, 27)


def xlsx_bytes(build) -> bytes:
    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def parse_csv_text(text: str, name="m.csv", encoding="utf-8", **kw):
    return mt.parse_metadata_table(text.encode(encoding), name, **kw)


def fixture(name: str) -> dict:
    return mt.parse_metadata_table((FIX / name).read_bytes(), name)


def has(msgs, needle):
    return any(needle in m for m in msgs)


def row_msgs(rows, key):
    return [m for r in rows for m in r[key]]


# --- XLSX ---------------------------------------------------------------------

def _typed(wb):
    ws = wb.active
    ws.title = "Muestras"
    ws.append(["ID muestra", "Organismo", "Hospedador", "País", "Fecha de muestreo", "Tecnología", "Profundidad",
               "Fórmula"])
    ws.append([1, "Escherichia coli", "Homo sapiens", "España", dt.date(2024, 3, 15), "MiSeq", 35.5, "=G2*2"])
    ws.append(["001", "Klebsiella pneumoniae", "Gallus gallus", "Perú", "2024-03-15", "MinION", 120, None])
    ws.append([])
    ws.append([2.0, "Salmonella enterica", None, "Zoë 北京 😀", "15/03/2024", "HiFi", 1e-3, None])
    ws.append([3, "S. aureus", None, None, dt.datetime(2024, 3, 15, 13, 45), "PacBio Sequel II", True, None])
    ws["A7"] = 7
    ws["A7"].number_format = "000"
    ws["B7"] = "E. faecium"
    ws["E7"] = 45366
    ws.merge_cells("G5:H5")  # merging C5:D5 would erase D5: openpyxl keeps only the top-left value
    wb.create_sheet("Notas").append(["x"])


def test_xlsx_types_and_conversions():
    p = mt.parse_metadata_table(xlsx_bytes(_typed), "meta.xlsx")
    assert p["kind"] == "xlsx" and p["sheet"] == "Muestras" and p["sheets"] == ["Muestras", "Notas"]
    rows = {r["row"]: r["values"] for r in p["rows"]}
    assert 4 not in rows
    assert rows[2]["ID muestra"] == "1"
    assert rows[3]["ID muestra"] == "001"
    assert rows[5]["ID muestra"] == "2"
    assert rows[7]["ID muestra"] == "007"
    assert has(row_msgs(p["rows"], "warnings"), "display format '000'")
    assert rows[2]["Fecha de muestreo"] == "2024-03-15"
    assert rows[6]["Fecha de muestreo"] == "2024-03-15T13:45:00"
    assert rows[5]["País"] == "Zoë 北京 😀"
    assert rows[5]["Profundidad"] == "0.001" and rows[6]["Profundidad"] == "TRUE"
    assert rows[2]["Fórmula"] is None
    assert has(p["warnings"], "1 formula cell(s) (H2)") and has(p["warnings"], "have no stored value")
    assert has(p["warnings"], "Merged cells (G5:H5)")
    assert has(p["warnings"], "The workbook has 2 sheets")
    assert p["numeric_cells"]["ID muestra"] == {2, 5, 6, 7}


def test_xlsx_cached_formula_values_libreoffice_fixture():
    p = fixture("lo_resaved_typed.xlsx")
    rows = {r["row"]: r["values"] for r in p["rows"]}
    assert rows[2]["Contig N50 (fórmula)"] == "71"
    assert has(p["warnings"], "formula cell(s)") and not has(p["warnings"], "no stored value")


def test_xlsx_build_samples_end_to_end():
    out = mt.build_samples(mt.parse_metadata_table(xlsx_bytes(_typed), "meta.xlsx"), today=TODAY)
    m = out["mapping"]
    assert m["ID muestra"] == "sample_code" and m["Hospedador"] == "host" and m["País"] == "geographic_origin"
    assert m["Fecha de muestreo"] == "collection_date" and m["Tecnología"] == "sequencing_technology"
    assert m["Organismo"] == "organism" and m["Profundidad"] == "custom"
    s = {x["row"]: x for x in out["rows"]}
    assert s[2]["data"]["sequencing_technology"] == "illumina"
    assert s[3]["data"]["sequencing_technology"] == "nanopore"
    assert s[5]["data"]["sequencing_technology"] == "pacbio"
    assert s[5]["data"]["collection_date"] == "2024-03-15"
    assert s[6]["data"]["collection_date"] == "2024-03-15"
    assert has(s[6]["warnings"], "time of day")
    assert s[2]["custom"] == {"Profundidad": "35.5"}
    assert s[7]["errors"] and has(s[7]["errors"], "Excel date serial")
    assert has(out["warnings"], "holds numbers in Excel")


def test_xlsx_header_below_title_rows_and_sheet_selection():
    def build(wb):
        wb.active.append(["Portada"])
        ws = wb.create_sheet("Datos")
        ws.append(["Proyecto vigilancia 2024"])
        ws.append([])
        ws.append(["Sample ID", "Host", "Country", "Collection date", "Platform"])
        ws.append(["S1", "human", "Spain", dt.date(2023, 1, 2), "Illumina NovaSeq 6000"])
    data = xlsx_bytes(build)
    with pytest.raises(mt.MetadataImportError, match="not found"):
        mt.parse_metadata_table(data, "m.xlsx", sheet="Nope")
    p = mt.parse_metadata_table(data, "m.xlsx", sheet="Datos")
    assert p["header_row"] == 3 and p["columns"][0] == "Sample ID"
    assert has(p["warnings"], "The header was detected on row 3; row(s) 1 above")
    out = mt.build_samples(p, today=TODAY)
    assert out["rows"][0]["data"] == {"sample_code": "S1", "host": "human", "geographic_origin": "Spain",
                                      "collection_date": "2023-01-02", "sequencing_technology": "illumina"}
    p2 = mt.parse_metadata_table(data, "m.xlsx", sheet="Datos", header_row=1)
    assert p2["columns"] == ["Proyecto vigilancia 2024", "column_B", "column_C", "column_D", "column_E"]
    with pytest.raises(mt.MetadataImportError, match="row 2 is empty"):
        mt.parse_metadata_table(data, "m.xlsx", sheet="Datos", header_row=2)


def test_xlsx_1904_epoch():
    def build(wb):
        wb.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
        wb.active.append(["sample", "date"])
        wb.active.append(["S1", dt.date(2024, 3, 15)])
    p = mt.parse_metadata_table(xlsx_bytes(build), "m.xlsx")
    assert p["rows"][0]["values"]["date"] == "2024-03-15" and has(p["warnings"], "1904")


def test_xlsx_wrong_dimension_is_ignored():
    data = xlsx_bytes(_typed)
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for i in src.infolist():
            b = src.read(i.filename)
            if i.filename == "xl/worksheets/sheet1.xml":
                # A regex: with lxml installed openpyxl writes '<dimension ref=".."/>' without the space.
                b, n = re.subn(rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="A1:A1"/>', b)
                assert n == 1
            z.writestr(i, b)
    p = mt.parse_metadata_table(out.getvalue(), "m.xlsx")
    assert len(p["columns"]) == 8 and len(p["rows"]) == 5


def test_xlsx_rejections():
    with pytest.raises(mt.MetadataImportError, match=r"\.xls files are not supported; save the sheet as \.xlsx or CSV UTF-8"):
        mt.parse_metadata_table(mt.OLE2_MAGIC + b"\0" * 100, "old.xls")
    with pytest.raises(mt.MetadataImportError, match="OLE2"):
        mt.parse_metadata_table(mt.OLE2_MAGIC + b"\0" * 100, "renamed.xlsx")
    with pytest.raises(mt.MetadataImportError, match=r"not an \.xlsx"):
        mt.parse_metadata_table(b"a,b\n1,2\n", "fake.xlsx")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(mt.MetadataImportError, match="not an Excel workbook"):
        mt.parse_metadata_table(buf.getvalue(), "zip.xlsx")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("[Content_Types].xml", "<x/>")
        z.writestr("xl/workbook.xml", "<x/>")
        z.writestr("xl/sharedStrings.xml", b"A" * (20 * 1024 * 1024))
    with pytest.raises(mt.MetadataTooLargeError, match="zip bomb"):
        mt.parse_metadata_table(buf.getvalue(), "bomb.xlsx")
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for i in range(mt.XLSX_MAX_MEMBERS + 1):
            z.writestr(f"p{i}", "")
    with pytest.raises(mt.MetadataTooLargeError, match="parts; the limit is 2000"):
        mt.parse_metadata_table(buf.getvalue(), "many.xlsx")
    for ext in (".ods", ".xlsm", ".xlsb"):
        with pytest.raises(mt.MetadataImportError, match=r"save the sheet as \.xlsx or CSV UTF-8"):
            mt.parse_metadata_table(b"PK\x03\x04", f"x{ext}")


def test_xlsx_entity_expansion_rejected():
    data = (FIX / "lo_resaved_typed.xlsx").read_bytes()
    lol = (b'<?xml version="1.0"?><!DOCTYPE l [<!ENTITY a "aaaaaaaaaa">'
           + b"".join(b'<!ENTITY a%d "%s">' % (i, (b"&a%d;" % (i - 1) if i > 1 else b"&a;") * 10) for i in range(1, 9))
           + b']><sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><si><t>&a8;</t></si></sst>')
    src, out = zipfile.ZipFile(io.BytesIO(data)), io.BytesIO()
    with zipfile.ZipFile(out, "w") as z:
        for i in src.infolist():
            z.writestr(i, lol if i.filename == "xl/sharedStrings.xml" else src.read(i.filename))
    with pytest.raises(mt.MetadataImportError, match="cannot open the workbook"):
        mt.parse_metadata_table(out.getvalue(), "lol.xlsx")



def _header_only(wb):
    wb.active.append(["sample_code", "host"])


def with_sheet_rows(extra_rows: str) -> bytes:
    """A small openpyxl workbook whose sheet gets raw <row> XML appended."""
    src, out = zipfile.ZipFile(io.BytesIO(xlsx_bytes(_header_only))), io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for i in src.infolist():
            b = src.read(i.filename)
            if i.filename == "xl/worksheets/sheet1.xml":
                b = b.replace(b"</sheetData>", extra_rows.encode() + b"</sheetData>")
            z.writestr(i, b)
    return out.getvalue()


def test_xlsx_style_only_cells_do_not_widen_rows(monkeypatch):
    calls = []
    real = mt._xlsx_cell_to_text
    monkeypatch.setattr(mt, "_xlsx_cell_to_text", lambda *a: calls.append(1) or real(*a))
    styled = "".join(f'<row r="{r}"><c r="XFD{r}" s="0"/></row>' for r in range(2, 2002))
    p = mt.parse_metadata_table(with_sheet_rows(styled), "w.xlsx")
    assert p["columns"] == ["sample_code", "host"] and p["rows"] == []
    assert len(calls) == 2  # only the header's two cells; empty formatted cells are never materialised


def test_xlsx_structure_is_bounded_before_reading(monkeypatch):
    with pytest.raises(mt.MetadataTooLargeError, match="cell XFD2 holds a value in column 16384; the limit is 256"):
        mt.parse_metadata_table(with_sheet_rows('<row r="2"><c r="XFD2"><v>1</v></c></row>'), "w.xlsx")
    for bad in ('<row r="999999999999"/>', '<row r="1"/>', '<row r="2.5"/>', '<row r="2"><c r="A?2"><v>1</v></c></row>'):
        with pytest.raises(mt.MetadataImportError, match="sheet 'Sheet'"):
            mt.parse_metadata_table(with_sheet_rows(bad), "w.xlsx")
    assert mt.parse_metadata_table(with_sheet_rows('<row r="9"><c r="A9" t="inlineStr"><is><t>S9</t></is></c></row>'),
                                   "w.xlsx")["rows"][0]["values"]["sample_code"] == "S9"
    many = "".join(f'<row r="{r}"/>' for r in range(2, 40))
    monkeypatch.setattr(mt, "XLSX_MAX_ROW_ELEMENTS", 30)
    with pytest.raises(mt.MetadataTooLargeError, match="more than 30 rows, counting empty formatted ones"):
        mt.parse_metadata_table(with_sheet_rows(many), "w.xlsx")
    monkeypatch.setattr(mt, "XLSX_MAX_CELL_ELEMENTS", 30)
    cells = "<row r=\"2\">" + "".join(f'<c r="{get_column_letter(i)}2" s="0"/>' for i in range(1, 40)) + "</row>"
    with pytest.raises(mt.MetadataTooLargeError, match="more than 30 cells"):
        mt.parse_metadata_table(with_sheet_rows(cells), "w.xlsx")
    monkeypatch.setattr(mt, "XLSX_MAX_GRID_CELLS", 1000)
    with pytest.raises(mt.MetadataTooLargeError, match="values reach row 900 and column 2 \\(B900\\)"):
        mt.parse_metadata_table(with_sheet_rows('<row r="900"><c r="B900"><v>1</v></c></row>'), "w.xlsx")
    data_rows = "".join(f'<row r="{r}"><c r="A{r}"><v>{r}</v></c></row>' for r in range(2, 40))
    with pytest.raises(mt.MetadataTooLargeError, match="more than 5 data rows"):
        mt.parse_metadata_table(with_sheet_rows(data_rows), "w.xlsx", limits=mt.Limits(max_rows=5))


def test_xlsx_shared_strings_are_bounded_before_openpyxl_loads_them(monkeypatch):
    data = (FIX / "lo_resaved_typed.xlsx").read_bytes()
    sst = zipfile.ZipFile(io.BytesIO(data)).read("xl/sharedStrings.xml")
    n = sst.count(b"<si>")
    assert n > 5 and mt.parse_metadata_table(data, "lo.xlsx")["rows"]
    monkeypatch.setattr(mt, "XLSX_MAX_SHARED_STRINGS", n - 1)
    with pytest.raises(mt.MetadataTooLargeError, match=f"more than {n - 1} distinct texts"):
        mt.parse_metadata_table(data, "lo.xlsx")
    monkeypatch.setattr(mt, "XLSX_MAX_PART_BYTES", len(sst) - 1)
    with pytest.raises(mt.MetadataTooLargeError, match=f"the limit is {len(sst) - 1:,}"):
        mt.parse_metadata_table(data, "lo.xlsx")


def test_xlsx_damaged_zip_directory_is_a_clear_error():
    data = bytearray(xlsx_bytes(_header_only))
    entry = data.find(b"PK\x01\x02")
    data[entry + 9] |= 0x08  # general-purpose flag bit 11: the name is UTF-8
    data[entry + 46:entry + 48] = b"\xc3\x28"  # but its bytes are not
    with pytest.raises(mt.MetadataImportError, match=r"not a valid \.xlsx \(zip\) file: UnicodeDecodeError"):
        mt.parse_metadata_table(bytes(data), "bad.xlsx")

def test_limits():
    small = mt.Limits(max_rows=20, max_columns=3, max_cell_chars=5)
    text = "id,host\n" + "".join(f"S{i},h\n" for i in range(30))
    with pytest.raises(mt.MetadataTooLargeError, match="more than 20 data rows"):
        parse_csv_text(text, limits=small)
    with pytest.raises(mt.MetadataTooLargeError, match="4 columns; the limit is 3"):
        parse_csv_text("a,b,c,d\n1,2,3,4\n", limits=small)
    with pytest.raises(mt.MetadataTooLargeError, match="characters; the limit is 5"):
        parse_csv_text("sample,notes\nS1,abcdefgh\n", limits=small)
    with pytest.raises(mt.MetadataImportError, match="empty"):
        mt.parse_metadata_table(b"", "m.csv")
    # trailing empty fields do not count as columns
    assert parse_csv_text("a,b\n1,2" + "," * 50 + "\n", limits=small)["columns"] == ["a", "b"]

    def wide(wb):
        wb.active.append(["sample"] + [f"c{i}" for i in range(10)])
    with pytest.raises(mt.MetadataTooLargeError, match="the limit is 5"):
        mt.parse_metadata_table(xlsx_bytes(wide), "w.xlsx", limits=mt.Limits(max_columns=5))

    def long_cell(wb):
        wb.active.append(["sample", "notes"])
        wb.active.append(["S1", "x" * 20])
    with pytest.raises(mt.MetadataTooLargeError, match="cell B2"):
        mt.parse_metadata_table(xlsx_bytes(long_cell), "l.xlsx", limits=mt.Limits(max_cell_chars=10))


def test_total_text_budget(monkeypatch):
    monkeypatch.setattr(mt, "MAX_TOTAL_CHARS", 100)
    with pytest.raises(mt.MetadataTooLargeError, match="more than 100 characters"):
        parse_csv_text("sample,notes\n" + "".join(f"S{i},{'x' * 20}\n" for i in range(10)))


# --- CSV ----------------------------------------------------------------------

def test_csv_bom_crlf_quoted_comma():
    p = parse_csv_text("﻿Sample ID,Host,Country,Date\r\nS1,human,Spain,2024-03-15\r\n"
                       "S2,\"chicken, broiler\",Peru,2024-03-16\r\n")
    assert p["encoding"] == "utf-8-sig" and p["delimiter"] == ","
    assert p["columns"] == ["Sample ID", "Host", "Country", "Date"]
    assert p["rows"][1]["values"]["Host"] == "chicken, broiler"


def test_csv_semicolon_spanish_cp1252_fixture():
    p = fixture("lo_export_cp1252_semicolon.csv")
    assert p["encoding"] == "cp1252" and p["delimiter"] == ";"
    assert has(p["warnings"], "Windows-1252")
    assert "País" in p["columns"] and "Tecnología" in p["columns"]
    rows = {r["row"]: r["values"] for r in p["rows"]}
    assert rows[2]["País"] == "España"
    assert rows[3]["Investigador"] == "Zoë ?? ?"  # characters outside cp1252 were lost by the export itself
    assert rows[9]["ID muestra"] == "007"  # LibreOffice exported the displayed (formatted) value
    out = mt.build_samples(p, today=TODAY)
    assert {s["row"]: s["sample_code"] for s in out["rows"]}[5] == "1.23456789012346E+019"


def test_csv_utf8_semicolon_fixture_keeps_unicode():
    p = fixture("lo_export_utf8_semicolon.csv")
    assert p["encoding"] == "utf-8" and not has(p["warnings"], "Windows-1252")
    assert {r["row"]: r["values"] for r in p["rows"]}[3]["Investigador"] == "Zoë 北京 😀"


def test_csv_utf8_semicolon_decimal_comma_kept_as_text():
    p = parse_csv_text("ID muestra;Hospedador;País;Profundidad\r\nM1;humano;España;35,5\r\n")
    assert p["delimiter"] == ";" and p["rows"][0]["values"]["Profundidad"] == "35,5"
    assert mt.build_samples(p, today=TODAY)["rows"][0]["custom"] == {"Profundidad": "35,5"}


def test_csv_tab_and_tsv_and_utf16():
    assert parse_csv_text("sample\thost\nS1\thuman\n")["delimiter"] == "\t"
    assert parse_csv_text("sample\thost\nS1\thuman\n", name="m.tsv")["delimiter"] == "\t"
    p = mt.parse_metadata_table("sample\thost\r\nS1\thuésped\r\n".encode("utf-16"), "m.txt")
    assert p["encoding"] == "utf-16" and p["rows"][0]["values"]["host"] == "huésped"
    assert parse_csv_text("sample|host\nS1|a|\n")["delimiter"] == "|"
    assert parse_csv_text("sample;host\nS1;a\n", delimiter=",")["columns"] == ["sample;host"]
    with pytest.raises(mt.MetadataImportError, match="delimiter must be one of"):
        parse_csv_text("a,b\n1,2\n", delimiter=",;")


def test_csv_embedded_newline_and_escaped_quotes():
    p = parse_csv_text('sample,notes\r\nS1,"line1\r\nline2"\r\nS2,"he said ""hi"""\r\n')
    assert [r["values"]["notes"] for r in p["rows"]] == ["line1\r\nline2", 'he said "hi"']
    assert [r["row"] for r in p["rows"]] == [2, 3]  # record numbers, not physical lines


def test_csv_single_column_and_sep_directive_and_ragged():
    p = parse_csv_text("sample\nS1\nS2\n")
    assert p["columns"] == ["sample"] and has(p["warnings"], "single column")
    p = parse_csv_text("sep=;\nsample;host\nS1;human\n")
    # Excel hides the sep= line, so the header is row 1 as in the spreadsheet view.
    assert p["delimiter"] == ";" and p["columns"] == ["sample", "host"] and p["header_row"] == 1
    assert p["rows"][0]["row"] == 2
    p = parse_csv_text("sample,host,country\nS1,human\nS2,cow,Chile,EXTRA\n")
    assert has(p["warnings"], "could not be detected reliably")
    assert p["rows"][1]["errors"] == ["the row has 4 fields but the header has 3; extra values ['EXTRA'] "
                                      "(an unquoted delimiter inside a field?)"]
    assert p["rows"][0]["values"]["country"] is None
    out = mt.build_samples(p, today=TODAY)
    assert [bool(s["errors"]) for s in out["rows"]] == [False, True]


def test_csv_ambiguous_delimiter_and_encoding_errors():
    with pytest.raises(mt.MetadataImportError, match="equally often"):
        parse_csv_text("a,b;c\n1,2;3,4\n1\n")
    with pytest.raises(mt.MetadataImportError, match="NUL"):
        mt.parse_metadata_table(b"a,b\n1,\x002\n", "m.csv")
    with pytest.raises(mt.MetadataImportError, match="neither valid UTF-8"):
        mt.parse_metadata_table(b"a,b\n\x81\x8d,1\n", "m.csv")
    with pytest.raises(mt.MetadataImportError, match="zip archive"):
        mt.parse_metadata_table(xlsx_bytes(_typed), "renamed.csv")
    with pytest.raises(mt.MetadataImportError, match="unsupported file type"):
        mt.parse_metadata_table(b"x", "m.json")


def test_csv_duplicate_and_missing_headers():
    assert parse_csv_text("a,a,a (2),A\n1,2,3,4\n")["columns"] == ["a", "a (2)", "a (2) (2)", "A (2)"]
    p = parse_csv_text("sample,Host,host,,notes\nS1,a,b,c,d\n")
    assert p["columns"] == ["sample", "Host", "host (2)", "column_D", "notes"]
    assert has(p["warnings"], "Duplicate header 'host' in column C") and has(p["warnings"], "no header")
    assert [m.code for m in p["warnings"]] == ["header_duplicate", "column_unnamed"]
    assert p["warnings"][0].params == {"header": "host", "column": "C", "name": "host (2)"}
    out = mt.build_samples(p, today=TODAY)
    assert out["mapping"]["host (2)"] == "custom" and has(out["warnings"], "both look like host")
    assert out["rows"][0]["custom"] == {"host (2)": "b", "column_D": "c"}


def test_control_characters_block_the_row_and_header():
    p = parse_csv_text("sample,host\nS\x0b1,a\nS2,b\n")
    assert has(p["rows"][0]["errors"], "control characters") and not p["rows"][1]["errors"]
    with pytest.raises(mt.MetadataImportError, match="control characters"):
        parse_csv_text("sam\x0bple,host\nS1,a\n")


# --- Mapping ------------------------------------------------------------------

@pytest.mark.parametrize("header,field", [
    ("Sample ID", "sample_code"), ("sample", "sample_code"), ("MUESTRA", "sample_code"),
    ("ID muestra", "sample_code"), ("ID_Muestra", "sample_code"), ("Código de muestra", "sample_code"),
    ("host", "host"), ("Hospedador", "host"), ("Huésped", "host"),
    ("country", "geographic_origin"), ("Origen", "geographic_origin"), ("Geographic origin", "geographic_origin"),
    ("País", "geographic_origin"), ("geo_loc_name", "geographic_origin"),
    ("date", "collection_date"), ("Fecha de muestreo", "collection_date"),
    ("Fecha de muestreo (dd/mm/aaaa)", "collection_date"),
    ("technology", "sequencing_technology"), ("Tecnología", "sequencing_technology"),
    ("TECNOLOGIA", "sequencing_technology"), ("Platform", "sequencing_technology"),
    ("Especie esperada", "expected_species"), ("Tipo de muestra", "sample_type"), ("Corrida", "sequencing_run"),
    ("Investigador", "researcher"), ("Institución", "institution"), ("Observaciones", "notes"),
    ("Descripción", "description"), ("Organismo", "organism"),
    ("MLST", "custom"), ("Resistencia", "custom"), ("ID", "custom"), ("Sample IDs", "custom"),
])
def test_header_mapping(header, field):
    assert mt.map_columns([header])["mapping"][header] == field
    assert mt.suggest_field(header) == field


def test_synonyms_do_not_clash():
    keys = [mt.normalize_label(s) for syns in mt.SYNONYMS.values() for s in syns]
    by_key: dict[str, set[str]] = {}
    for field, syns in mt.SYNONYMS.items():
        for s in syns:
            by_key.setdefault(mt.normalize_label(s), set()).add(field)
    assert all(len(v) == 1 for v in by_key.values()) and len(by_key) <= len(keys)


def test_mapping_overrides():
    r = mt.map_columns(["Cepa", "Lugar", "x"], overrides={"Cepa": "sample_code", "Lugar": "geographic_origin",
                                                          "x": "ignore"})
    assert r["mapping"] == {"Cepa": "sample_code", "Lugar": "geographic_origin", "x": "ignore"} and not r["warnings"]
    # an explicit choice takes the field even when a later column would have been suggested for it
    r = mt.map_columns(["host", "Hospedador"], overrides={"Hospedador": "host"})
    assert r["mapping"] == {"host": "custom", "Hospedador": "host"} and has(r["warnings"], "both look like host")
    with pytest.raises(mt.MetadataImportError, match="unknown target"):
        mt.map_columns(["x"], overrides={"x": "not_a_field"})
    with pytest.raises(mt.MetadataImportError, match="not in the file"):
        mt.map_columns(["x"], overrides={"y": "host"})
    with pytest.raises(mt.MetadataImportError, match="both mapped to host"):
        mt.map_columns(["a", "b"], overrides={"a": "host", "b": "host"})


def test_missing_sample_code_column_blocks_every_row():
    out = mt.build_samples(parse_csv_text("Cepa,host\nC1,a\n"), today=TODAY)
    assert has(out["errors"], "No column is mapped to sample_code")
    assert out["rows"][0]["errors"] == ["no column is mapped to sample_code"]
    ok = mt.build_samples(parse_csv_text("Cepa,host\nC1,a\n"), overrides={"Cepa": "sample_code"}, today=TODAY)
    assert not ok["errors"] and ok["rows"][0]["data"] == {"sample_code": "C1", "host": "a"}


def test_ignored_columns_are_not_imported():
    out = mt.build_samples(parse_csv_text("sample,MLST,secret\nS1,ST131,x\n"), overrides={"secret": "ignore"},
                           today=TODAY)
    assert out["rows"][0]["custom"] == {"MLST": "ST131"}


@pytest.mark.parametrize("value,expected,warns", [
    ("Illumina", "illumina", False), ("MiSeq", "illumina", False), ("NovaSeq 6000", "illumina", False),
    ("Illumina NextSeq 500", "illumina", False), ("iSeq100", "illumina", False),
    ("ONT", "nanopore", False), ("MinION Mk1B", "nanopore", False), ("Oxford Nanopore", "nanopore", False),
    ("nanopore", "nanopore", False), ("PromethION", "nanopore", False),
    ("HiFi", "pacbio", False), ("PacBio Sequel II", "pacbio", False), ("Revio", "pacbio", False),
    ("Híbrido", "hybrid", False), ("Illumina + ONT", "hybrid", True), ("MiSeq/MinION", "hybrid", True),
    ("Ion Torrent", "unknown", True), ("DNBSEQ-G400", "unknown", True), ("Sanger", "unknown", True),
    ("foo", "unknown", True), ("front", "unknown", True), ("montana", "unknown", True),
    ("", None, False), (None, None, False), ("N/A", None, False), ("desconocido", None, False),
    ("unknown", None, False),
])
def test_technology(value, expected, warns):
    tech, msg = mt.normalize_technology(value)
    assert tech == expected and (msg is not None) == warns


def test_technology_original_kept_when_unrecognised():
    out = mt.build_samples(parse_csv_text("sample,platform\nS1,Ion Torrent S5\nS2,MiSeq\nS3,N/A\n"), today=TODAY)
    s1, s2, s3 = out["rows"]
    assert "sequencing_technology" not in s1["data"] and s1["custom"] == {"platform (original)": "Ion Torrent S5"}
    assert has(s1["warnings"], "kept as text")
    assert s2["data"]["sequencing_technology"] == "illumina" and not s2["custom"]
    assert "sequencing_technology" not in s3["data"] and not s3["custom"] and not s3["warnings"]


# --- Dates --------------------------------------------------------------------

@pytest.mark.parametrize("text,order,status,iso", [
    ("2024-03-15", None, "ok", "2024-03-15"),
    ("2024-03-15T10:00:00", None, "ok", "2024-03-15"),
    ("2024/03/15", None, "ok", "2024-03-15"),
    ("15/03/2024", None, "ok", "2024-03-15"),
    ("03/15/2024", None, "ok", "2024-03-15"),
    ("15-03-2024", None, "ok", "2024-03-15"),
    ("15.03.2024", None, "ok", "2024-03-15"),
    ("05/05/2024", None, "ok", "2024-05-05"),
    ("03/05/2024", None, "ambiguous", None),
    ("03/05/2024", "DMY", "ok", "2024-05-03"),
    ("03/05/2024", "MDY", "ok", "2024-03-05"),
    ("15/03/2024", "MDY", "invalid", None),
    ("03/15/2024", "DMY", "invalid", None),
    ("31/02/2024", None, "invalid", None),
    ("2024-02-30", None, "invalid", None),
    ("15/03/24", None, "two_digit_year", None),
    ("2024", None, "partial", None),
    ("2024-03", None, "partial", None),
    ("15 de marzo de 2024", None, "ok", "2024-03-15"),
    ("March 15, 2024", None, "ok", "2024-03-15"),
    ("15-Mar-2024", None, "ok", "2024-03-15"),
    ("1 ago 2023", None, "ok", "2023-08-01"),
    ("45366", None, "serial", None),
    ("yesterday", None, "unrecognised", None),
    ("٢٠٢٤-٠٣-١٥", None, "unrecognised", None),
    ("", None, "empty", None), ("N/A", None, "empty", None), ("not collected", None, "empty", None),
])
def test_parse_date(text, order, status, iso):
    r = mt.parse_date(text, order)
    assert r.status == status
    assert (r.value.isoformat() if r.value else None) == iso


def test_date_order_never_inferred_from_other_rows():
    p = parse_csv_text("sample;fecha\nA;15/03/2024\nB;03/05/2024\n")
    out = mt.build_samples(p, today=TODAY)
    s = {x["sample_code"]: x for x in out["rows"]}
    assert s["A"]["data"]["collection_date"] == "2024-03-15"
    assert s["B"]["errors"] and "collection_date" not in s["B"]["data"]
    assert s["B"]["custom"] == {"fecha (original)": "03/05/2024"}
    assert out["ambiguous_dates"] == 1 and out["date_order"] is None
    assert has(out["warnings"], "Ambiguous dates are not assumed")
    out2 = mt.build_samples(p, date_order="DMY", today=TODAY)
    assert {x["sample_code"]: x["data"].get("collection_date") for x in out2["rows"]}["B"] == "2024-05-03"
    assert out2["ambiguous_dates"] == 0 and out2["date_order_source"] == "user"


def test_date_order_from_header_hint_and_mixed_columns():
    out = mt.build_samples(parse_csv_text("sample,Fecha de muestreo (dd/mm/aaaa)\nA,03/05/2024\n"), today=TODAY)
    assert out["date_order"] == "DMY" and out["date_order_source"] == "header"
    assert out["rows"][0]["data"]["collection_date"] == "2024-05-03"
    out = mt.build_samples(parse_csv_text("sample,date\nA,15/03/2024\nB,03/15/2024\n"), today=TODAY)
    assert has(out["warnings"], "mixes day-first")


def test_partial_and_future_dates_warn_without_blocking():
    out = mt.build_samples(parse_csv_text("sample,date\nB,2030-01-01\nC,2024\n"), today=TODAY)
    b, c = out["rows"]
    assert not b["errors"] and has(b["warnings"], "in the future")
    assert not c["errors"] and has(c["warnings"], "partial date") and c["custom"] == {"date (original)": "2024"}
    assert "collection_date" not in c["data"]


# --- Sample codes -------------------------------------------------------------

def test_sample_code_rules():
    out = mt.build_samples(parse_csv_text("sample,host\nS1,a\n,c\n" + "X" * 300 + ",e\n\"S\n2\",f\n"), today=TODAY)
    msgs = row_msgs(out["rows"], "errors")
    assert "sample_code is empty" in msgs
    assert any("limit is 255" in m for m in msgs) and not any(m == "sample_code is empty" for m in out["rows"][2]["errors"])
    assert "sample_code contains a tab or line break" in msgs
    assert [bool(s["errors"]) for s in out["rows"]] == [False, True, True, True]


def test_unicode_normalisation_nfc():
    p = parse_csv_text("sample,researcher\nS1,José\n")
    assert p["rows"][0]["values"]["researcher"] == "José"


def test_float_sample_code_from_excel_is_rejected():
    def build(wb):
        wb.active.append(["sample", "host"])
        wb.active.append([12345678901234567890, "h"])
        wb.active.append([1.5, "h"])
        wb.active.append([42, "h"])
    out = mt.build_samples(mt.parse_metadata_table(xlsx_bytes(build), "m.xlsx"), today=TODAY)
    assert [s["sample_code"] for s in out["rows"]] == ["1.234567890123457e+19", "1.5", "42"]
    assert [bool(s["errors"]) for s in out["rows"]] == [True, True, False]
    assert has(row_msgs(out["rows"], "errors"), "cannot be recovered")


def test_scientific_notation_and_formula_prefix_warnings_csv():
    out = mt.build_samples(parse_csv_text("sample;host\n1.23456789012346E+019;h\n=cmd;h\n"), today=TODAY)
    assert not row_msgs(out["rows"], "errors")
    warnings = row_msgs(out["rows"], "warnings")
    assert has(warnings, "scientific notation") and has(warnings, "formula character")


def test_openpyxl_warnings_are_forwarded_and_error_cells_empty():
    def build(wb):
        ws = wb.active
        ws.append(["sample", "date"])
        ws.append(["S1", 99999999])
        ws["B2"].number_format = "yyyy-mm-dd"
        ws.append(["S2", "#N/A"])
    p = mt.parse_metadata_table(xlsx_bytes(build), "m.xlsx")
    assert [r["values"]["date"] for r in p["rows"]] == [None, None]
    assert has(p["warnings"], "openpyxl: Cell B2 is marked as a date but the serial value 99999999")
    assert has(p["rows"][1]["warnings"], "cell B3: Excel error value '#N/A'")


# --- Planning against existing samples ---------------------------------------

def _existing(code, **fields):
    return {"id": f"id-{code}", "sample_code": code, "host": None, "custom_metadata": None, **fields}


def test_plan_create_update_unchanged_and_errors():
    built = mt.build_samples(parse_csv_text("sample,host,MLST\nNEW,cow,\ns1,human,ST1\nS2,pig,\nX,,\nx,,\n"),
                             today=TODAY)
    existing = [_existing("S1", host="human"), _existing("S2", host="pig"), _existing("X")]
    plan = {r["row"]: r for r in mt.plan_rows(built, existing)}
    assert plan[2]["action"] == "create" and plan[2]["changes"][0] == {"field": "sample_code", "key": None,
                                                                       "old": None, "new": "NEW"}
    assert plan[3]["action"] == "update" and plan[3]["matched_by"] == "case_insensitive"
    assert plan[3]["matched_sample_code"] == "S1" and plan[3]["sample_id"] == "id-S1"
    assert plan[3]["changes"] == [{"field": "custom_metadata", "key": "MLST", "old": None, "new": "ST1"}]
    assert plan[4]["action"] == "unchanged" and plan[4]["matched_by"] == "exact"
    assert plan[5]["action"] == "unchanged"
    assert plan[6]["action"] == "error" and has(plan[6]["errors"], "also appears on row 5 as 'X'")


def test_plan_ambiguous_case_insensitive_match_is_an_error():
    built = mt.build_samples(parse_csv_text("sample,host\nab,cow\n"), today=TODAY)
    plan = mt.plan_rows(built, [_existing("AB"), _existing("Ab")])
    assert plan[0]["action"] == "error" and has(plan[0]["errors"], "matches several existing samples")


def test_plan_compares_custom_values_as_text():
    built = mt.build_samples(parse_csv_text("sample,depth,mlst\nS1,35,ST1\nS2,36,ST1\n"), today=TODAY)
    existing = [_existing(c, custom_metadata={"depth": 35, "mlst": "ST1"}) for c in ("S1", "S2")]
    plan = mt.plan_rows(built, existing)
    assert plan[0]["action"] == "unchanged"
    assert plan[1]["changes"] == [{"field": "custom_metadata", "key": "depth", "old": "35", "new": "36"}]


# --- Export -------------------------------------------------------------------

@pytest.mark.parametrize("value,expected", [
    ("=1+2", "'=1+2"), ("+34 600", "'+34 600"), ("-80C", "'-80C"), ("@SUM(A1)", "'@SUM(A1)"),
    ("\t=x", "'\t=x"), ("\r=x", "'\r=x"), ("\n=x", "'\n=x"), ("＝cmd", "'＝cmd"), ("＠x", "'＠x"),
    ("normal", "normal"), ("E. coli", "E. coli"), (None, ""),
    ("-34.6", "-34.6"), ("+12", "+12"), ("-1e5", "-1e5"), ("-34,6", "-34,6"), ("-", "'-"), ("=", "'="),
    ("'=x", "''=x"), ("''-", "'''-"), ("'-34.6", "'-34.6"), ("'note", "'note"), ("'", "'"),
])
def test_sanitize(value, expected):
    assert mt.sanitize_spreadsheet_cell(value) == expected


@pytest.mark.parametrize("value", ["=1+2", "-", "+", "'=x", "''+y", "'-34.6", "-34.6", "'", "''", "'a", "x", "＝x", "\t=x"])
def test_csv_neutralisation_is_undone_exactly_once(value):
    exported = mt.sanitize_spreadsheet_cell(value)
    assert mt.is_neutralised_cell(exported) == (exported != value)
    assert (exported[1:] if mt.is_neutralised_cell(exported) else exported) == value


def test_export_csv_roundtrip_and_injection():
    samples = [{"sample_code": "S1", "notes": '=HYPERLINK("http://x","y")', "sequencing_technology": "illumina",
                "custom_metadata": {"MLST": "ST131", "evil": "@cmd|' /C calc'!A0", "lat": "-34.6", "n": 3}},
               {"sample_code": "-S2", "custom_metadata": {"=x": "y"}}]
    data = mt.export_csv(samples)
    assert data.startswith(b"\xef\xbb\xbf") and b"\r\n" in data
    rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"), newline="")))
    assert rows[0][len(mt.ARVA_FIELDS):] == ["MLST", "evil", "lat", "n", "'=x"]
    assert rows[1][mt.ARVA_FIELDS.index("notes")] == '\'=HYPERLINK("http://x","y")'
    assert rows[1][-5:] == ["ST131", "'@cmd|' /C calc'!A0", "-34.6", "3", ""]
    assert rows[2][0] == "'-S2"
    first_line = data.decode("utf-8-sig").split("\r\n")[0]
    assert first_line.startswith('"sample_code","organism"')  # QUOTE_ALL


def test_xlsx_export_cells_are_text():
    samples = [{"sample_code": "=1+1", "host": "#N/A", "notes": "a\x0bb", "sequencing_technology": "pacbio"}]
    wb = openpyxl.load_workbook(io.BytesIO(mt.export_xlsx(samples)), read_only=False, data_only=False)
    ws = wb.active
    assert ws["A2"].value == "=1+1" and ws["A2"].data_type == "s"
    assert ws["E2"].value == "#N/A" and ws["E2"].data_type == "s"
    assert ws.cell(2, mt.ARVA_FIELDS.index("notes") + 1).value == "a�b"
    assert ws.column_dimensions["A"].number_format == "@"
    assert [c.value for c in ws[1]][:2] == ["sample_code", "organism"]


I18N = Path(__file__).resolve().parents[2] / "web" / "src" / "i18n" / "index.ts"


@pytest.mark.skipif(not I18N.exists(), reason="the web sources are not next to the API")
def test_every_message_code_has_the_same_english_text_in_the_interface():
    from app.tasks.pod5_reader import WARNINGS

    source = I18N.read_text(encoding="utf-8")
    en = source[: source.index("const es:")]
    for prefix, registry in (("metadata.msg", mt.MESSAGES), ("pod5.msg", WARNINGS)):
        for code, template in registry.items():
            line = f"  {json.dumps(f'{prefix}.{code}')}: {json.dumps(template, ensure_ascii=False)},\n"
            assert line in en, f"{prefix}.{code} is missing or differs in apps/web/src/i18n/index.ts"
    assert mt.Msg("date_ambiguous", value="'1/2/2024'").params == {"value": "'1/2/2024'"}
    assert mt.message_json("free text") == {"code": None, "params": {}, "text": "free text"}
